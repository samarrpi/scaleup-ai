from __future__ import annotations

from io import BytesIO
import json
from pathlib import Path

import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from agents.process_understanding import understand_process
from agents.document_ingestion import understand_pdf, attach_document_context
from agents.scientific.scaleup import calculate_first_pass_scale_up
from agents.scientific.scaleup_engineering import (
    calculate_engineering_scale_up,
    suggest_engineering_inputs,
)
from agents.scientific.scaleup_assessment import assess_integrated_scale_up
from agents.scientific.model_registry import inspect_models
from agents.scientific.ecoli_runner import (
    get_ecoli_model_readiness,
    is_ecoli_model_applicable,
    run_from_process_state,
)
from schemas.process import ModelResult, Parameter, ProcessState, Provenance


ROOT = Path(__file__).resolve().parents[1]

EXAMPLES = {
    "Clean process": "clean_process.txt",
    "E. coli fed-batch": "ecoli_fed_batch.txt",
    "Ambiguous information": "ambiguous_information.txt",
    "Contradiction": "contradiction.txt",
    "Missing information": "missing_information.txt",
    "Messy process": "messy_process.txt",
}


def load_example(filename: str) -> str:
    """Return the selected bundled example without interpreting it."""
    return (ROOT / "examples" / filename).read_text(encoding="utf-8")


def load_selected_example() -> None:
    """Place the selected example in the editable description field."""
    st.session_state["process_description"] = load_example(
        EXAMPLES[st.session_state["example_name"]]
    )
    st.session_state["process_state"] = None
    st.session_state["model_result"] = None
    st.session_state["model_error"] = None
    st.session_state["model_applicable"] = False
    st.session_state["model_readiness"] = None
    st.session_state["scaleup_result"] = None
    st.session_state["engineering_scaleup_result"] = None
    st.session_state["phase3_assessment"] = None
    st.session_state["engineering_scaleup_basis_initialized"] = False
    st.session_state["engineering_scaleup_suggestion"] = None
    st.session_state["engineering_scaleup_basis_initialized"] = False
    st.session_state["engineering_scaleup_suggestion"] = None


def sanitize_excel_value(value: object) -> str:
    """Prevent malicious spreadsheet formulas from being exported as-is."""
    if value is None:
        return ""

    text = str(value).strip()

    if not text:
        return ""

    dangerous_starts = ("=", "+", "-", "@", ":", "cmd(", "CMD(", "java", "http", "file")
    if text[:1] in ("=", "+", "-", "@") or text.lower().startswith("cmd("):
        return "'" + text

    if text.lower().startswith(("=cmd", "+cmd", "-cmd", "@cmd")):
        return "'" + text

    return text


def create_excel_export(
    state: ProcessState,
    result: ModelResult | None,
    scaleup_result: dict | None = None,
    engineering_scaleup_result: dict | None = None,
    phase3_assessment: dict | None = None,
) -> bytes:
    """
    Create a single-sheet Excel export containing the
    current ProcessState and scientific model results.
    """

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "ScaleUp AI Analysis"

    headers = [
        "Section",
        "Item",
        "Value / Detail",
        "Units",
        "Status",
        "Notes",
    ]

    worksheet.append(headers)

    rows: list[list[str]] = []

    # ---------------------------------------------------------------
    # Process information
    # ---------------------------------------------------------------

    rows.extend(
        [
            [
                "Process",
                "Process",
                sanitize_excel_value(state.process or "Unknown"),
                "",
                "Extracted",
                "",
            ],
            [
                "Process",
                "Organism",
                sanitize_excel_value(state.organism or "Unknown"),
                "",
                "Extracted",
                "",
            ],
            [
                "Process",
                "Product",
                sanitize_excel_value(state.product or "Unknown"),
                "",
                "Extracted",
                "",
            ],
            [
                "Process",
                "Equipment",
                sanitize_excel_value(state.equipment or "Unknown"),
                "",
                "Extracted",
                "",
            ],
            [
                "Process",
                "Scale",
                sanitize_excel_value(state.scale or "Unknown"),
                "",
                "Extracted",
                "",
            ],
        ]
    )

    # ---------------------------------------------------------------
    # Parameters
    # ---------------------------------------------------------------

    for parameter in state.parameters:
        rows.append(
            [
                "Parameters",
                parameter.name,
                sanitize_excel_value(
                    str(parameter.value) if parameter.value is not None else ""
                ),
                parameter.unit or "",
                parameter.source.value,
                sanitize_excel_value(parameter.description or ""),
            ]
        )

    # ---------------------------------------------------------------
    # Objectives
    # ---------------------------------------------------------------

    for objective in state.objectives:
        rows.append(
            [
                "Objectives",
                objective.name,
                sanitize_excel_value(
                    str(objective.target) if objective.target is not None else ""
                ),
                "",
                objective.source.value,
                f"Direction: {objective.direction}",
            ]
        )

    # ---------------------------------------------------------------
    # Constraints
    # ---------------------------------------------------------------

    for constraint in state.constraints:
        rows.append(
            [
                "Constraints",
                constraint.name,
                sanitize_excel_value(
                    str(constraint.value) if constraint.value is not None else ""
                ),
                "",
                constraint.source.value,
                f"Condition: {constraint.condition}",
            ]
        )

    # ---------------------------------------------------------------
    # Quality checks
    # ---------------------------------------------------------------

    for finding in state.missing_information:
        rows.append(
            [
                "Quality checks",
                "Missing information",
                sanitize_excel_value(finding),
                "",
                "Flagged",
                "",
            ]
        )

    for finding in state.ambiguities:
        rows.append(
            [
                "Quality checks",
                "Ambiguity",
                sanitize_excel_value(finding),
                "",
                "Flagged",
                "",
            ]
        )

    for finding in state.contradictions:
        rows.append(
            [
                "Quality checks",
                "Contradiction",
                sanitize_excel_value(finding),
                "",
                "Flagged",
                "",
            ]
        )

    # ---------------------------------------------------------------
    # Evidence
    # ---------------------------------------------------------------

    for evidence in state.evidence:
        rows.append(
            [
                "Evidence",
                evidence.id,
                sanitize_excel_value(evidence.excerpt or ""),
                "",
                evidence.status.value,
                sanitize_excel_value(
                    f"Source: {evidence.source}; Location: {evidence.location or ''}; "
                    f"Page: {evidence.page or ''}; Limitation: {evidence.limitation or ''}"
                ),
            ]
        )

    # ---------------------------------------------------------------
    # Process graph
    # ---------------------------------------------------------------

    for node in state.process_nodes:
        rows.append(
            [
                "Process graph",
                node.name,
                node.kind,
                "",
                node.provenance.value,
                sanitize_excel_value(f"Node ID: {node.id}; Evidence: {', '.join(node.evidence_ids)}"),
            ]
        )

    for relation in state.process_relations:
        rows.append(
            [
                "Process graph",
                relation.relation,
                f"{relation.source_node_id} -> {relation.target_node_id}",
                "",
                relation.provenance.value,
                sanitize_excel_value(f"Evidence: {', '.join(relation.evidence_ids)}"),
            ]
        )

    # ---------------------------------------------------------------
    # Structured uncertainty
    # ---------------------------------------------------------------

    for finding in state.uncertainty_findings:
        rows.append(
            [
                "Uncertainty",
                finding.category,
                sanitize_excel_value(finding.message),
                "",
                "Flagged",
                sanitize_excel_value(f"Subject: {finding.subject or ''}; Evidence: {', '.join(finding.evidence_ids)}"),
            ]
        )

    # ---------------------------------------------------------------
    # Scientific model
    # ---------------------------------------------------------------

    if result is not None:
        model_metadata = result.model_metadata
        predictions = result.predictions

        rows.extend(
            [
                [
                    "Scientific model",
                    "Model name",
                    sanitize_excel_value(
                        model_metadata.get("model_name", "")
                    ),
                    "",
                    "Model",
                    "",
                ],
                [
                    "Scientific model",
                    "Model version",
                    sanitize_excel_value(
                        model_metadata.get("model_version", "")
                    ),
                    "",
                    "Model",
                    "",
                ],
                [
                    "Scientific model",
                    "Numerical method",
                    sanitize_excel_value(
                        model_metadata.get("numerical_method", "")
                    ),
                    "",
                    "Model",
                    "",
                ],
                [
                    "Scientific model",
                    "Time step",
                    sanitize_excel_value(
                        model_metadata.get("time_step_h", "")
                    ),
                    "h",
                    "Model",
                    "",
                ],
                [
                    "Scientific model",
                    "Final biomass",
                    sanitize_excel_value(
                        predictions.get("final_biomass_g_L", "")
                    ),
                    "g/L",
                    "Predicted",
                    "",
                ],
                [
                    "Scientific model",
                    "Final substrate",
                    sanitize_excel_value(
                        predictions.get("final_substrate_g_L", "")
                    ),
                    "g/L",
                    "Predicted",
                    "",
                ],
                [
                    "Scientific model",
                    "Final product",
                    sanitize_excel_value(
                        predictions.get("final_product_g_L", "")
                    ),
                    "g/L",
                    "Predicted",
                    "",
                ],
                [
                    "Scientific model",
                    "Final volume",
                    sanitize_excel_value(
                        predictions.get("final_volume_L", "")
                    ),
                    "L",
                    "Predicted",
                    "",
                ],
                [
                    "Scientific model",
                    "Simulation duration",
                    sanitize_excel_value(
                        predictions.get("simulation_time_h", "")
                    ),
                    "h",
                    "Predicted",
                    "",
                ],
            ]
        )

        # -----------------------------------------------------------
        # Input validation
        # -----------------------------------------------------------

        validation = model_metadata.get("input_validation", {})

        for parameter_name in validation.get("provided_parameters", []):
            rows.append(
                [
                    "Scientific input validation",
                    "Provided parameter",
                    sanitize_excel_value(parameter_name),
                    "",
                    "Provided",
                    "",
                ]
            )

        for parameter_name in validation.get("assumed_parameters", []):
            rows.append(
                [
                    "Scientific input validation",
                    "Assumed parameter",
                    sanitize_excel_value(parameter_name),
                    "",
                    "Assumption",
                    "",
                ]
            )

        for parameter_name in validation.get("missing_parameters", []):
            rows.append(
                [
                    "Scientific input validation",
                    "Missing parameter",
                    sanitize_excel_value(parameter_name),
                    "",
                    "Missing",
                    "",
                ]
            )

        # -----------------------------------------------------------
        # Assumption values
        # -----------------------------------------------------------

        for assumption in model_metadata.get("input_assumptions", []):
            rows.append(
                [
                    "Input assumptions",
                    assumption["parameter"],
                    sanitize_excel_value(str(assumption["value"])),
                    "",
                    "Assumed",
                    sanitize_excel_value(assumption["reason"]),
                ]
            )

        # -----------------------------------------------------------
        # Important parameters
        # -----------------------------------------------------------

        for parameter_name in result.important_parameters:
            rows.append(
                [
                    "Scientific model",
                    "Important parameter",
                    sanitize_excel_value(parameter_name),
                    "",
                    "Model parameter",
                    "",
                ]
            )

        # -----------------------------------------------------------
        # Uncertainty
        # -----------------------------------------------------------

        for key, value in result.uncertainty.items():
            rows.append(
                [
                    "Uncertainty",
                    str(key),
                    sanitize_excel_value(str(value)),
                    "",
                    "Model",
                    "",
                ]
            )

                # -----------------------------------------------------------
        # Engineering analysis
        # -----------------------------------------------------------

        engineering = model_metadata.get(
            "engineering_analysis",
            {},
        )

        if engineering.get("status") == "complete":
            engineering_rows = [
                (
                    "Duration",
                    engineering.get("duration_h"),
                    "h",
                ),
                (
                    "Initial volume",
                    engineering.get("initial_volume_L"),
                    "L",
                ),
                (
                    "Final volume",
                    engineering.get("final_volume_L"),
                    "L",
                ),
                (
                    "Feed volume",
                    engineering.get("feed_volume_L"),
                    "L",
                ),
                (
                    "Initial biomass",
                    engineering.get("initial_biomass_g_L"),
                    "g/L",
                ),
                (
                    "Final biomass",
                    engineering.get("final_biomass_g_L"),
                    "g/L",
                ),
                (
                    "Initial biomass mass",
                    engineering.get("initial_biomass_mass_g"),
                    "g",
                ),
                (
                    "Final biomass mass",
                    engineering.get("final_biomass_mass_g"),
                    "g",
                ),
                (
                    "Biomass formed",
                    engineering.get("biomass_formed_g"),
                    "g",
                ),
                (
                    "Initial substrate",
                    engineering.get("initial_substrate_g_L"),
                    "g/L",
                ),
                (
                    "Final substrate",
                    engineering.get("final_substrate_g_L"),
                    "g/L",
                ),
                (
                    "Initial substrate mass",
                    engineering.get("initial_substrate_mass_g"),
                    "g",
                ),
                (
                    "Final substrate mass",
                    engineering.get("final_substrate_mass_g"),
                    "g",
                ),
                (
                    "Substrate fed",
                    engineering.get("substrate_fed_g"),
                    "g",
                ),
                (
                    "Substrate consumed",
                    engineering.get("substrate_consumed_g"),
                    "g",
                ),
                (
                    "Initial product",
                    engineering.get("initial_product_g_L"),
                    "g/L",
                ),
                (
                    "Final product",
                    engineering.get("final_product_g_L"),
                    "g/L",
                ),
                (
                    "Initial product mass",
                    engineering.get("initial_product_mass_g"),
                    "g",
                ),
                (
                    "Final product mass",
                    engineering.get("final_product_mass_g"),
                    "g",
                ),
                (
                    "Product formed",
                    engineering.get("product_formed_g"),
                    "g",
                ),
                (
                    "Biomass yield",
                    engineering.get("biomass_yield_g_g"),
                    "g/g",
                ),
                (
                    "Product yield",
                    engineering.get("product_yield_g_g"),
                    "g/g",
                ),
                (
                    "Biomass productivity",
                    engineering.get("biomass_productivity_g_L_h"),
                    "g/L/h",
                ),
                (
                    "Product productivity",
                    engineering.get("product_productivity_g_L_h"),
                    "g/L/h",
                ),
                (
                    "Feed rate",
                    engineering.get("feed_rate_L_h"),
                    "L/h",
                ),
                (
                    "Feed substrate concentration",
                    engineering.get("feed_substrate_g_L"),
                    "g/L",
                ),
                (
                    "Dilution rate",
                    engineering.get("dilution_rate_1_h"),
                    "1/h",
                ),
                (
                    "Expected feed volume",
                    engineering.get("expected_feed_volume_L"),
                    "L",
                ),
                (
                    "Volume balance error",
                    engineering.get("volume_balance_error_L"),
                    "L",
                ),
                (
                    "Volume balance status",
                    engineering.get("volume_balance_status"),
                    "",
                ),
            ]

            for metric_name, value, unit in engineering_rows:
                rows.append(
                    [
                        "Engineering analysis",
                        metric_name,
                        sanitize_excel_value(
                            str(value) if value is not None else ""
                        ),
                        unit,
                        "Calculated",
                        "",
                    ]
                )

            for warning in engineering.get("warnings", []):
                rows.append(
                    [
                        "Engineering analysis",
                        "Warning",
                        sanitize_excel_value(str(warning)),
                        "",
                        "Warning",
                        "",
                    ]
                )

        # -----------------------------------------------------------
        # Simulation trajectory
        # -----------------------------------------------------------

        time_values = predictions.get("time_h", [])
        biomass_values = predictions.get("biomass_g_L", [])
        substrate_values = predictions.get("substrate_g_L", [])
        product_values = predictions.get("product_g_L", [])
        volume_values = predictions.get("volume_L", [])

        for index, time_value in enumerate(time_values):
            rows.append(
                [
                    "Simulation trajectory",
                    f"t = {time_value}",
                    sanitize_excel_value(
                        f"Biomass={biomass_values[index]:.6g}; "
                        f"Substrate={substrate_values[index]:.6g}; "
                        f"Product={product_values[index]:.6g}; "
                        f"Volume={volume_values[index]:.6g}"
                    ),
                    "h / g/L / g/L / g/L / L",
                    "Predicted",
                    "",
                ]
            )

        if scaleup_result is not None:
            scaleup_rows = [
                ("Scale-up basis", scaleup_result.get("scale_up_basis"), "", "Basis"),
                ("Laboratory working volume", scaleup_result.get("lab_volume_L"), "L", "Input"),
                ("Target initial working volume", scaleup_result.get("target_volume_L"), "L", "Input"),
                ("Maximum allowable working volume", scaleup_result.get("max_working_volume_L"), "L", "Input"),
                ("Scale factor", scaleup_result.get("scale_factor"), "×", "Calculated"),
                ("Large scale factor threshold", scaleup_result.get("large_scale_factor_threshold"), "×", "Screening threshold"),
                ("Large scale factor flag", scaleup_result.get("large_scale_factor_flag"), "", "Screening flag"),
                ("Laboratory feed rate", scaleup_result.get("feed_rate_L_h"), "L/h", "Input"),
                ("Scaled feed rate", scaleup_result.get("scaled_feed_rate_L_h"), "L/h", "Calculated"),
                ("Process duration", scaleup_result.get("duration_h"), "h", "Input"),
                ("Scaled feed volume", scaleup_result.get("scaled_feed_volume_L"), "L", "Calculated"),
                ("Expected final volume", scaleup_result.get("expected_final_volume_L"), "L", "Calculated"),
                ("Initial biomass mass", scaleup_result.get("scaled_initial_biomass_mass_g"), "g", "Calculated"),
                ("Initial substrate mass", scaleup_result.get("scaled_initial_substrate_mass_g"), "g", "Calculated"),
                ("Initial product mass", scaleup_result.get("scaled_initial_product_mass_g"), "g", "Calculated"),
                ("Feed substrate requirement", scaleup_result.get("scaled_substrate_feed_mass_g"), "g", "Calculated"),
                ("Vessel capacity status", scaleup_result.get("vessel_capacity_status"), "", "Engineering check"),
                ("Vessel capacity message", scaleup_result.get("vessel_capacity_message"), "", "Engineering check"),
            ]

            for metric_name, value, unit, status in scaleup_rows:
                rows.append(
                    [
                        "Scale-up analysis",
                        metric_name,
                        sanitize_excel_value(str(value) if value is not None else ""),
                        unit,
                        status,
                        "First-pass volumetric scaling",
                    ]
                )

            for check_name, check_status in scaleup_result.get("engineering_checks", {}).items():
                rows.append(
                    [
                        "Scale-up engineering checks",
                        check_name,
                        sanitize_excel_value(str(check_status)),
                        "",
                        "Assessment",
                        "Separate scale-dependent engineering analysis required where status is not evaluated.",
                    ]
                )

            for warning in scaleup_result.get("warnings", []):
                rows.append(
                    [
                        "Scale-up analysis",
                        "Warning",
                        sanitize_excel_value(str(warning)),
                        "",
                        "Warning",
                        "",
                    ]
                )

            if engineering_scaleup_result is not None:
                rows.append(
                    [
                        "Engineering scale-up",
                        "Basis",
                        sanitize_excel_value(engineering_scaleup_result.get("basis", "")),
                        "",
                        "Engineering screening",
                        "",
                    ]
                )

                geometry = engineering_scaleup_result.get("geometry", {})
                for name, value in geometry.items():
                    rows.append([
                        "Engineering geometry",
                        name,
                        sanitize_excel_value(str(value) if value is not None else ""),
                        "",
                        "Calculated",
                        "",
                    ])

                for section_name, section in [
                    ("Tip speed", engineering_scaleup_result.get("tip_speed", {})),
                    ("Power", engineering_scaleup_result.get("power", {})),
                    ("Aeration", engineering_scaleup_result.get("aeration", {})),
                    ("Oxygen transfer", engineering_scaleup_result.get("oxygen", {})),
                    ("Oxygen demand", engineering_scaleup_result.get("oxygen_demand", {})),
                    ("Heat transfer", engineering_scaleup_result.get("heat", {})),
                ]:
                    for name, value in section.items():
                        rows.append([
                            f"Engineering {section_name}",
                            name,
                            sanitize_excel_value(str(value) if value is not None else ""),
                            "",
                            "Calculated",
                            "",
                        ])

                for check_name, check_status in engineering_scaleup_result.get("checks", {}).items():
                    rows.append([
                        "Engineering assessment",
                        check_name,
                        sanitize_excel_value(str(check_status)),
                        "",
                        "Assessment",
                        "Screening result based only on supplied engineering inputs.",
                    ])

                for warning in engineering_scaleup_result.get("warnings", []):
                    rows.append([
                        "Engineering assessment",
                        "Warning",
                        sanitize_excel_value(str(warning)),
                        "",
                        "Warning",
                        "",
                    ])

        # -----------------------------------------------------------
        # Phase 3 integrated assessment
        # -----------------------------------------------------------

        if phase3_assessment is not None:
            rows.append([
                "Phase 3 assessment",
                "Assessment type",
                sanitize_excel_value(phase3_assessment.get("assessment_type", "")),
                "",
                "Assessment",
                "",
            ])
            rows.append([
                "Phase 3 assessment",
                "Scale-up basis",
                sanitize_excel_value(phase3_assessment.get("scale_up_basis", "")),
                "",
                "Basis",
                "",
            ])
            rows.append([
                "Phase 3 assessment",
                "Scale factor",
                sanitize_excel_value(str(phase3_assessment.get("scale_factor", ""))),
                "×",
                "Calculated",
                "",
            ])
            rows.append([
                "Phase 3 assessment",
                "Large scale transition",
                sanitize_excel_value(str(phase3_assessment.get("large_scale_transition", ""))),
                "",
                "Flag",
                "",
            ])

            completeness = phase3_assessment.get("data_completeness", {})
            for key, value in completeness.items():
                rows.append([
                    "Phase 3 data completeness",
                    key,
                    sanitize_excel_value(str(value)),
                    "",
                    "Calculated",
                    "",
                ])

            for criterion in phase3_assessment.get("criteria", []):
                rows.append([
                    "Phase 3 criteria",
                    criterion.get("criterion", ""),
                    sanitize_excel_value(str(criterion.get("target") if criterion.get("target") is not None else "")),
                    criterion.get("unit", ""),
                    criterion.get("status", ""),
                    sanitize_excel_value(
                        f"Category: {criterion.get('category', '')}; "
                        f"Basis: {criterion.get('basis', '')}; "
                        f"Lab: {criterion.get('laboratory')}; "
                        f"Ratio/utilisation: {criterion.get('ratio')}; "
                        f"Utilisation %: {criterion.get('utilisation_percent')}; "
                        f"Limit: {criterion.get('limit')}; "
                        f"Deviation %: {criterion.get('deviation_percent')}"
                    ),
                ])

            for requirement in phase3_assessment.get("data_requirements", []):
                rows.append([
                    "Phase 3 data requirements",
                    requirement.get("criterion", ""),
                    sanitize_excel_value(str(requirement.get("required_data", ""))),
                    "",
                    "Required",
                    "",
                ])

            for limitation in phase3_assessment.get("limitations", []):
                rows.append([
                    "Phase 3 limitations",
                    "Limitation",
                    sanitize_excel_value(str(limitation)),
                    "",
                    "Review",
                    "",
                ])

    else:
        rows.append(
            [
                "Scientific model",
                "Status",
                "No scientific model result available",
                "",
                "Not available",
                "",
            ]
        )

    for row in rows:
        worksheet.append(row)

    # ---------------------------------------------------------------
    # Formatting
    # ---------------------------------------------------------------

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    column_widths = {
        1: 28,
        2: 34,
        3: 72,
        4: 25,
        5: 22,
        6: 55,
    }

    for column_index, width in column_widths.items():
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    worksheet.auto_filter.ref = worksheet.dimensions

    if worksheet.max_row >= 2:
        table = Table(
            displayName="ScaleUpAIAnalysis",
            ref=f"A1:F{worksheet.max_row}",
        )

        table_style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )

        table.tableStyleInfo = table_style
        worksheet.add_table(table)

    worksheet.row_dimensions[1].height = 30

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()

def display_source_explanation(state: ProcessState) -> None:
    """Explain which source currently drives the ProcessState and how to interpret it."""
    document = (state.metadata or {}).get("document_source")
    supporting = (state.metadata or {}).get("supporting_document")
    if document:
        st.info(
            f"**Source basis:** {document}. ScaleUp AI extracted page-aware text from this PDF. "
            "Claims are only promoted into the ProcessState when the current extraction rules support them."
        )
    elif supporting:
        st.info(
            f"**Primary source:** process description. **Supporting document:** {supporting.get('source', 'PDF')} "
            f"({supporting.get('page_count', 0)} pages). The PDF is retained as document context; it does not overwrite the primary process description."
        )
    else:
        st.info(
            "**Source basis:** user-provided process description. Values shown here are extracted from that text; "
            "ScaleUp AI does not silently fill missing scientific information."
        )


def display_document_context(state: ProcessState) -> None:
    """Show document coverage and transparent keyword signals when a PDF is present."""
    info = (state.metadata or {}).get("supporting_document")
    if not info:
        info = {
            "source": (state.metadata or {}).get("document_source"),
            "document_type": (state.metadata or {}).get("document_type"),
            "page_count": (state.metadata or {}).get("page_count"),
            "page_text_extracted": (state.metadata or {}).get("page_text_extracted"),
            "keyword_signals": (state.metadata or {}).get("keyword_signals", {}),
            "role": "primary_document",
        } if (state.metadata or {}).get("document_type") == "pdf" else None
    if not info:
        return

    st.markdown("### Document coverage")
    st.caption(
        "This is document-level information, not a claim about the process. Keyword signals are shown as discovery aids only and are not automatically treated as scientific facts."
    )
    c1, c2, c3 = st.columns(3)
    c1.metric("Pages", info.get("page_count") or 0)
    c2.metric("Pages with text", info.get("page_text_extracted") or 0)
    signals = info.get("keyword_signals") or {}
    c3.metric("Keyword signals", len(signals))
    if signals:
        st.dataframe(
            [{"Term": key, "Mentions": value} for key, value in signals.items()],
            width="stretch", hide_index=True,
        )
        st.caption("A keyword signal means the term occurs in the document. It does not establish that the term is the process identity or an operating condition.")



def display_document_intelligence(state: ProcessState) -> None:
    """Explain the current document-level extraction without overstating it as fact."""
    report = (state.metadata or {}).get("document_intelligence")
    if not report:
        supporting = (state.metadata or {}).get("supporting_document") or {}
        report = supporting.get("document_intelligence")
    if not report:
        return

    st.subheader("Document intelligence")
    st.info(
        "**What this means:** ScaleUp AI first observes the document, then separates "
        "reviewable candidate claims from validated process facts. A candidate claim is "
        "not automatically promoted into the ProcessState just because a keyword appears."
    )

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Pages", report.get("page_count", 0))
    c2.metric("Pages with text", report.get("pages_with_text", 0))
    c3.metric("Sections", len(report.get("sections", [])))
    c4.metric("Candidate claims", len(report.get("claims", [])))

    warnings = report.get("warnings", [])
    if warnings:
        with st.expander("Document limitations and warnings", expanded=False):
            for warning in warnings:
                st.warning(warning)

    sections = report.get("sections", [])
    if sections:
        st.markdown("#### Detected document structure")
        st.caption("These are coarse structural observations used to organise the paper; they are not scientific conclusions.")
        st.dataframe(
            [
                {
                    "Section": item.get("title", ""),
                    "Type": item.get("section_type", "other"),
                    "Pages": f"{item.get('start_page')}–{item.get('end_page')}",
                    "Preview": item.get("text_preview", ""),
                }
                for item in sections
            ],
            width="stretch",
            hide_index=True,
        )

    claims = report.get("claims", [])
    if claims:
        st.markdown("#### Candidate scientific claims")
        st.caption(
            "These observations were found by transparent rules. They retain page and section context but remain **candidate** until a semantic review layer validates them."
        )
        st.dataframe(
            [
                {
                    "Claim ID": claim.get("id"),
                    "Type": claim.get("claim_type"),
                    "Page": claim.get("page") or "—",
                    "Section": claim.get("section") or "—",
                    "Candidate": claim.get("text", ""),
                    "Status": claim.get("promotion_status", "candidate"),
                }
                for claim in claims
            ],
            width="stretch",
            hide_index=True,
        )
    else:
        st.info("No rule-based candidate claims were detected. This is not evidence that the paper contains no scientific information; it means the current conservative rules did not promote any observations to candidates.")


def display_process_state(state: ProcessState) -> None:
    """Display the structured process knowledge without downstream analysis."""

    summary = st.columns(5)
    for column, label, value in zip(
        summary,
        ["Process", "Organism", "Product", "Equipment", "Scale"],
        [state.process, state.organism, state.product, state.equipment, state.scale],
    ):
        column.metric(label, value or "Unknown")

    display_source_explanation(state)
    display_document_context(state)
    display_document_intelligence(state)

    st.subheader("Parameters")
    if state.parameters:
        st.dataframe(
            [
                {
                    "Name": parameter.name,
                    "Value": parameter.value,
                    "Unit": parameter.unit or "",
                    "Type": getattr(parameter, "parameter_type", ""),
                    "Provenance": parameter.source.value,
                    "Evidence": ", ".join(parameter.evidence_ids) or "—",
                }
                for parameter in state.parameters
            ],
            width="stretch",
            hide_index=True,
        )
    else:
        st.info("No parameters were identified.")

    objectives, constraints = st.columns(2)
    with objectives:
        st.subheader("Objectives")
        if state.objectives:
            st.dataframe(
                [
                    {
                        "Name": objective.name,
                        "Direction": objective.direction,
                        "Target": objective.target,
                        "Provenance": objective.source.value,
                        "Evidence": ", ".join(objective.evidence_ids) or "—",
                    }
                    for objective in state.objectives
                ],
                width="stretch",
                hide_index=True,
            )
        else:
            st.info("No objectives were identified.")

    with constraints:
        st.subheader("Constraints")
        if state.constraints:
            st.dataframe(
                [
                    {
                        "Name": constraint.name,
                        "Condition": constraint.condition,
                        "Value": constraint.value,
                        "Provenance": constraint.source.value,
                        "Evidence": ", ".join(constraint.evidence_ids) or "—",
                    }
                    for constraint in state.constraints
                ],
                width="stretch",
                hide_index=True,
            )
        else:
            st.info("No constraints were identified.")


def display_process_graph(state: ProcessState) -> None:
    """Render a compact, inspectable and evidence-linked process topology."""
    st.subheader("Process graph")

    st.info(
        "**What this graph is for:** it shows the process structure that ScaleUp AI can "
        "support from the supplied source. Each box is a process entity and each arrow "
        "is a source-supported relationship. The graph is intentionally conservative: "
        "it will not invent material flows, causal links, or operating conditions that "
        "were not reported."
    )

    if not state.process_nodes:
        st.info(
            "No explicit operation sequence was extracted. ScaleUp AI will not invent "
            "process topology from a bare process label. Add an 'Operations:' or "
            "'Process steps:' sequence to the description to build the graph."
        )
        return

    node_lookup = {node.id: node for node in state.process_nodes}
    evidence_lookup = {item.id: item for item in state.evidence}

    relation_explanations = {
        "precedes": "Sequence only: the source places the first entity before the second.",
        "produces": "The source explicitly states that the source entity produces the target.",
        "consumes": "The source explicitly states that the source entity consumes the target.",
        "transfers_to": "The source explicitly states that material or state is transferred to the target.",
    }

    operations = [node for node in state.process_nodes if node.kind == "operation"]
    material_states = [node for node in state.process_nodes if node.kind == "material_state"]
    relation_types = sorted({relation.relation for relation in state.process_relations})

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Nodes", len(state.process_nodes))
    c2.metric("Relations", len(state.process_relations))
    c3.metric("Operations", len(operations))
    c4.metric("Relation types", len(relation_types))

    with st.expander("How to read this", expanded=False):
        st.markdown(
            """
            **Operation** = a process step explicitly reported by the source.

            **Material state** = a material or state explicitly represented in the structured process.

            **precedes** = sequence information only. It does **not** mean mass transfer, causation,
            or that one step consumes the other.

            **produces / consumes / transfers_to** = material relationships, shown only when the
            source supports them.

            **Evidence** = the source record supporting the node or relationship. Select a node
            below to inspect the evidence and any linked uncertainty.
            """
        )

    # Compact left-to-right topology. Detailed information lives in the inspector below.
    lines = [
        "digraph G {",
        "rankdir=LR;",
        'graph [pad="0.10", nodesep="0.22", ranksep="0.35", bgcolor="transparent"];',
        'node [shape=box, style="rounded,filled", fontname="Arial", fontsize=9, margin="0.10,0.06", height=0.42];',
        'edge [fontname="Arial", fontsize=7, color="#9CA3AF", arrowsize=0.55];',
    ]

    fill_by_kind = {
        "operation": "#1F2937",
        "material_state": "#243B53",
    }

    for node in state.process_nodes:
        escaped_name = node.name.replace("\\", "\\\\").replace('"', '\\"')
        kind_label = "Operation" if node.kind == "operation" else node.kind.replace("_", " ").title()
        evidence_count = len(node.evidence_ids)
        label = f"{kind_label}\\n{escaped_name}\\nEvidence: {evidence_count}"
        fill = fill_by_kind.get(node.kind, "#202124")
        lines.append(
            f'"{node.id}" [label="{label}", fillcolor="{fill}", fontcolor="white", color="#6B7280"];'
        )

    for relation in state.process_relations:
        left = node_lookup.get(relation.source_node_id)
        right = node_lookup.get(relation.target_node_id)
        if left and right:
            edge_label = relation.relation.replace('"', '\\"')
            lines.append(f'"{left.id}" -> "{right.id}" [label="{edge_label}"];')

    lines.append("}")
    graph_col, _ = st.columns([2.15, 0.85])
    with graph_col:
        st.graphviz_chart("\n".join(lines), use_container_width=True)

    st.caption(
        "The diagram is deliberately compact. Use the inspector below for the detailed "
        "meaning, evidence, and uncertainty behind a node."
    )

    st.markdown("### Inspect a process node")
    node_options = [node.id for node in state.process_nodes]
    selected_id = st.selectbox(
        "Select a node",
        node_options,
        format_func=lambda node_id: (
            f"{node_lookup[node_id].name} — {node_lookup[node_id].kind.replace('_', ' ').title()}"
        ),
        key="process_graph_node_inspector",
    )
    selected_node = node_lookup[selected_id]

    left, right = st.columns([1, 1])
    with left:
        st.markdown(f"**{selected_node.name}**")
        st.write(f"**Type:** {selected_node.kind.replace('_', ' ').title()}")
        st.write(f"**Node ID:** `{selected_node.id}`")
        st.write(f"**Provenance:** `{selected_node.provenance.value}`")
        if selected_node.description:
            st.write(f"**Description:** {selected_node.description}")
        else:
            st.caption("No additional node description was explicitly reported.")

        linked_relations = [
            relation
            for relation in state.process_relations
            if relation.source_node_id == selected_node.id
            or relation.target_node_id == selected_node.id
        ]
        st.markdown("**Connected relationships**")
        if linked_relations:
            for relation in linked_relations:
                source_name = node_lookup[relation.source_node_id].name
                target_name = node_lookup[relation.target_node_id].name
                st.write(f"- `{source_name}` → **{relation.relation}** → `{target_name}`")
                st.caption(relation_explanations.get(relation.relation, "Structured relationship retained without additional interpretation."))
        else:
            st.info("No connected relationships are recorded for this node.")

    with right:
        st.markdown("**Node evidence**")
        node_evidence = [
            evidence_lookup[evidence_id]
            for evidence_id in selected_node.evidence_ids
            if evidence_id in evidence_lookup
        ]
        if node_evidence:
            for evidence in node_evidence:
                with st.expander(f"{evidence.id} — {evidence.status.value}", expanded=False):
                    st.write(f"**Source:** {evidence.source}")
                    if evidence.page is not None:
                        st.write(f"**Page:** {evidence.page}")
                    if evidence.location:
                        st.write(f"**Location:** {evidence.location}")
                    if evidence.excerpt:
                        st.code(evidence.excerpt, language="text")
                    if evidence.limitation:
                        st.write(f"**Limitation:** {evidence.limitation}")
        else:
            st.info("No evidence record is attached to this node.")

        linked_uncertainty = [
            finding
            for finding in state.uncertainty_findings
            if any(evidence_id in selected_node.evidence_ids for evidence_id in finding.evidence_ids)
        ]
        if linked_uncertainty:
            st.markdown("**Uncertainty linked to this evidence**")
            for finding in linked_uncertainty:
                st.warning(f"{finding.category.title()}: {finding.message}")

    st.markdown("### Reported process conditions")
    st.caption(
        "These are process-level facts. They are not assigned to an individual operation unless "
        "the source explicitly provides that stage association."
    )
    if state.parameters:
        st.dataframe(
            [
                {
                    "Parameter": parameter.name,
                    "Value": parameter.value,
                    "Unit": parameter.unit or "—",
                    "Provenance": parameter.source.value,
                    "Evidence": ", ".join(parameter.evidence_ids) or "—",
                }
                for parameter in state.parameters
            ],
            width="stretch",
            hide_index=True,
        )
    else:
        st.info("No process parameters were extracted.")

    with st.expander("Relationship audit", expanded=False):
        st.dataframe(
            [
                {
                    "Source": node_lookup[r.source_node_id].name,
                    "Relationship": r.relation,
                    "Target": node_lookup[r.target_node_id].name,
                    "Meaning": relation_explanations.get(r.relation, "Structured relationship retained without additional interpretation."),
                    "Provenance": r.provenance.value,
                    "Evidence": ", ".join(r.evidence_ids) or "—",
                }
                for r in state.process_relations
                if r.source_node_id in node_lookup and r.target_node_id in node_lookup
            ],
            width="stretch",
            hide_index=True,
        )

    with st.expander("What this graph currently does not claim"):
        st.markdown(
            "- It does **not** infer unreported material flows.\n"
            "- It does **not** assign a parameter to an operation when the source does not provide that stage association.\n"
            "- It does **not** convert a sequence relationship into a mass, energy, or causal relationship.\n"
            "- It does **not** resolve contradictory source statements; those remain in the Uncertainty layer."
        )


def display_evidence(state: ProcessState) -> None:
    """Show the claim → source → location → excerpt evidence chain."""
    st.subheader("Evidence & provenance")
    st.info("**What this means:** Evidence connects a structured claim to the source location that supports it. A PDF page is a source location; a keyword signal alone is not evidence for a claim.")
    if not state.evidence:
        st.info("No source evidence records are available for this ProcessState.")
        return

    st.caption(
        "For text input, evidence points to the user-provided process description. For PDFs, "
        "validated evidence can retain the source page. Document-level candidate claims are shown "
        "separately so that keyword observations are not mistaken for validated facts."
    )
    st.dataframe(
        [
            {
                "Evidence ID": evidence.id,
                "Source": evidence.source,
                "Location": evidence.location or "—",
                "Page": evidence.page or "—",
                "Excerpt": evidence.excerpt or "—",
                "Status": evidence.status.value,
                "Limitation": evidence.limitation or "—",
            }
            for evidence in state.evidence
        ],
        width="stretch",
        hide_index=True,
    )


def display_uncertainty(state: ProcessState) -> None:
    """Display missing, ambiguous, and contradictory findings as structured records."""
    st.subheader("Uncertainty & missing information")
    st.info("**What this means:** Missing = not reported; Ambiguous = reported but unclear; Contradictory = conflicting statements. These findings are kept visible instead of being silently resolved.")
    counts = {
        "Reported": len(state.parameters) + len(state.objectives) + len(state.constraints),
        "Missing": len(state.missing_information),
        "Ambiguous": len(state.ambiguities),
        "Contradictory": len(state.contradictions),
    }
    cols = st.columns(4)
    for col, (label, count) in zip(cols, counts.items()):
        col.metric(label, count)

    if state.uncertainty_findings:
        st.dataframe(
            [
                {
                    "Category": finding.category.title(),
                    "Message": finding.message,
                    "Subject": finding.subject or "—",
                    "Evidence": ", ".join(finding.evidence_ids) or "—",
                }
                for finding in state.uncertainty_findings
            ],
            width="stretch",
            hide_index=True,
        )
    else:
        st.success("No missing, ambiguous, or contradictory findings were recorded.")


def display_model_registry(state: ProcessState) -> None:
    """Show applicable scientific models before any simulation is run."""
    st.subheader("Scientific model registry")
    st.info("**What this means:** Applicable asks whether a registered model matches this process. Ready asks whether that model has enough process-specific inputs to run. Neither status is a scientific validation of the process.")
    records = inspect_models(state)
    for record in records:
        with st.container(border=True):
            left, right = st.columns([3, 1])
            with left:
                st.markdown(f"**{record['name']}**")
                st.caption(record["description"])
            with right:
                st.metric("Applicable", "Yes" if record["applicable"] else "No")
                if record["applicable"]:
                    st.metric("Ready", "Yes" if record["ready"] else "No")
                    readiness = record.get("readiness") or {}
                    missing = readiness.get("missing_required_parameters", [])
                    if missing:
                        st.caption("Missing model inputs: " + ", ".join(missing))
                    elif record["ready"]:
                        st.caption("All required process-specific inputs are currently available.")
                else:
                    st.caption("This model is not currently supported by the extracted process identity.")


def display_simulation_status(
    state: ProcessState,
    model_applicable: bool,
    model_readiness: dict | None,
    model_result: ModelResult | None,
    model_error: str | None,
) -> None:
    """Render the model-readiness gate and simulation output."""
    st.info("**What this means:** Simulation executes a registered scientific model using the current ProcessState. If inputs are missing, ScaleUp AI asks you to supply them rather than inventing them.")

    if not model_applicable:
        st.info("No currently registered scientific model is applicable to this ProcessState.")
        return

    st.success("E. coli fed-batch model is applicable to this ProcessState.")
    if model_readiness and model_readiness.get("ready"):
        st.success("Model readiness: all required process-specific inputs are available.")
    elif model_readiness:
        st.warning("Model applicability is established, but required process-specific inputs are missing.")
        missing_parameters = model_readiness.get("missing_required_parameters", [])
        if missing_parameters:
            st.markdown("**Missing process inputs**")
            for name in missing_parameters:
                st.write(f"- {name}")
        display_ecoli_input_form(state, missing_parameters)

    if model_error:
        st.error(f"Scientific model error: {model_error}")

    if model_result:
        display_model_result(model_result, state)


def display_input_validation(result: ModelResult, state: ProcessState) -> None:
    """Display scientific input provenance and validation."""
    validation = result.model_metadata.get("input_validation", {})
    assumptions = result.model_metadata.get("input_assumptions", [])

    if not validation:
        return

    st.subheader("Scientific input validation")

    provided_names = validation.get("provided_parameters", [])
    assumed_names = validation.get("assumed_parameters", [])
    missing_names = validation.get("missing_parameters", [])
    status = validation.get("status", "unknown")

    # ---------------------------------------------------------------
    # Build lookup of parameters supplied by the process description
    # ---------------------------------------------------------------

    parameter_lookup = {}
    for parameter in state.parameters:
        parameter_lookup[parameter.name.strip().lower()] = parameter

    # ---------------------------------------------------------------
    # Process-provided inputs
    # ---------------------------------------------------------------

    provided_rows = []
    for parameter_name in provided_names:
        parameter = parameter_lookup.get(parameter_name.strip().lower())
        if parameter is not None:
            provided_rows.append(
                {
                    "Parameter": parameter_name,
                    "Value": parameter.value,
                    "Unit": parameter.unit or "",
                }
            )

    # ---------------------------------------------------------------
    # Prototype assumptions
    # ---------------------------------------------------------------

    assumption_lookup = {
        assumption["parameter"]: assumption
        for assumption in assumptions
    }

    assumption_units = {
        "initial_product": "g/L",
        "mu_max": "1/h",
        "Ks": "g/L",
        "Yxs": "g/g",
        "Ypx": "g/g",
        "initial_biomass": "g/L",
        "initial_substrate": "g/L",
        "initial_volume": "L",
        "feed_rate": "L/h",
        "feed_substrate": "g/L",
        "duration": "h",
    }

    assumed_rows = []
    for parameter_name in assumed_names:
        assumption = assumption_lookup.get(parameter_name)
        if assumption is not None:
            assumed_rows.append(
                {
                    "Parameter": parameter_name,
                    "Value": assumption["value"],
                    "Unit": assumption_units.get(parameter_name, ""),
                }
            )

    # ---------------------------------------------------------------
    # Validation status
    # ---------------------------------------------------------------

    if status == "complete" and assumed_rows:
        st.warning(
            f"The scientific model used {len(provided_rows)} process-provided input(s) "
            f"and {len(assumed_rows)} prototype assumption(s)."
        )
    elif status == "incomplete":
        st.error("The scientific model is missing required inputs.")
    else:
        st.success("All required scientific inputs were provided.")

    # ---------------------------------------------------------------
    # Process-provided inputs table
    # ---------------------------------------------------------------

    if provided_rows:
        st.markdown("### Process-provided inputs")
        st.dataframe(provided_rows, width="stretch", hide_index=True)

    # ---------------------------------------------------------------
    # Prototype assumptions table
    # ---------------------------------------------------------------

    if assumed_rows:
        st.markdown("### Model assumptions")
        st.dataframe(assumed_rows, width="stretch", hide_index=True)

    # ---------------------------------------------------------------
    # Missing required inputs
    # ---------------------------------------------------------------

    if missing_names:
        st.markdown("### Missing required inputs")
        for parameter_name in missing_names:
            st.write(f"- {parameter_name}")


def display_model_trajectories(result: ModelResult) -> None:
    """Display scientific model trajectories."""
    predictions = result.predictions

    time_values = predictions.get("time_h", [])
    biomass_values = predictions.get("biomass_g_L", [])
    substrate_values = predictions.get("substrate_g_L", [])
    product_values = predictions.get("product_g_L", [])
    volume_values = predictions.get("volume_L", [])

    if not time_values:
        st.info("No simulation trajectory is available.")
        return

    st.subheader("Process trajectories")
    st.caption(
        "These plots show the simulated evolution of the process states over time."
    )

    biomass_chart = {
        "Time (h)": time_values,
        "Biomass (g/L)": biomass_values,
    }
    st.line_chart(biomass_chart, x="Time (h)", y="Biomass (g/L)", width="stretch")

    substrate_chart = {
        "Time (h)": time_values,
        "Substrate (g/L)": substrate_values,
    }
    st.line_chart(substrate_chart, x="Time (h)", y="Substrate (g/L)", width="stretch")

    product_chart = {
        "Time (h)": time_values,
        "Product (g/L)": product_values,
    }
    st.line_chart(product_chart, x="Time (h)", y="Product (g/L)", width="stretch")

    volume_chart = {
        "Time (h)": time_values,
        "Volume (L)": volume_values,
    }
    st.line_chart(volume_chart, x="Time (h)", y="Volume (L)", width="stretch")

def display_model_result(
    result: ModelResult,
    state: ProcessState,
) -> None:
    """Display the scientific model output."""

    display_input_validation(
        result,
        state,
    )

    st.subheader("Scientific model")

    predictions = result.predictions

    metrics = st.columns(4)

    metrics[0].metric(
        "Final biomass",
        f"{predictions['final_biomass_g_L']:.2f} g/L",
    )

    metrics[1].metric(
        "Final substrate",
        f"{predictions['final_substrate_g_L']:.2f} g/L",
    )

    metrics[2].metric(
        "Final product",
        f"{predictions['final_product_g_L']:.2f} g/L",
    )

    metrics[3].metric(
        "Final volume",
        f"{predictions['final_volume_L']:.2f} L",
    )

    model_name = result.model_metadata.get(
        "model_name",
        "Unknown",
    )

    st.caption(
        f"Model: {model_name}"
    )

    # ---------------------------------------------------------------
    # Engineering analysis
    # ---------------------------------------------------------------

    engineering = result.model_metadata.get(
        "engineering_analysis",
        {},
    )

    if engineering.get("status") == "complete":
        st.subheader("Engineering analysis")

        engineering_metrics = st.columns(4)

        product_yield = engineering.get(
            "product_yield_g_g"
        )

        productivity = engineering.get(
    "product_productivity_g_L_h"
)

        substrate_consumed = engineering.get(
            "substrate_consumed_g"
        )

        feed_volume = engineering.get(
            "feed_volume_L"
        )

        engineering_metrics[0].metric(
            "Product yield",
            (
                f"{product_yield:.3f} g/g"
                if product_yield is not None
                else "N/A"
            ),
        )

        engineering_metrics[1].metric(
            "Productivity",
            (
                f"{productivity:.3f} g/L/h"
                if productivity is not None
                else "N/A"
            ),
        )

        engineering_metrics[2].metric(
            "Substrate consumed",
            (
                f"{substrate_consumed:.2f} g"
                if substrate_consumed is not None
                else "N/A"
            ),
        )

        engineering_metrics[3].metric(
            "Feed volume",
            (
                f"{feed_volume:.2f} L"
                if feed_volume is not None
                else "N/A"
            ),
        )

        with st.expander(
            "Engineering analysis details"
        ):
            analysis_rows = [
                {
                    "Metric": "Initial biomass",
                    "Value": engineering.get(
                        "initial_biomass_g_L"
                    ),
                    "Unit": "g/L",
                },
                {
                    "Metric": "Final biomass",
                    "Value": engineering.get(
                        "final_biomass_g_L"
                    ),
                    "Unit": "g/L",
                },
                {
                    "Metric": "Biomass formed",
                    "Value": engineering.get(
                        "biomass_formed_g"
                    ),
                    "Unit": "g",
                },
                {
                    "Metric": "Initial substrate",
                    "Value": engineering.get(
                        "initial_substrate_g_L"
                    ),
                    "Unit": "g/L",
                },
                {
                    "Metric": "Final substrate",
                    "Value": engineering.get(
                        "final_substrate_g_L"
                    ),
                    "Unit": "g/L",
                },
                {
                    "Metric": "Substrate fed",
                    "Value": engineering.get(
                        "substrate_fed_g"
                    ),
                    "Unit": "g",
                },
                {
                    "Metric": "Substrate consumed",
                    "Value": engineering.get(
                        "substrate_consumed_g"
                    ),
                    "Unit": "g",
                },
                {
                    "Metric": "Biomass yield",
                    "Value": engineering.get(
                        "biomass_yield_g_g"
                    ),
                    "Unit": "g/g",
                },
                {
                    "Metric": "Product yield",
                    "Value": engineering.get(
                        "product_yield_g_g"
                    ),
                    "Unit": "g/g",
                },
                {
    "Metric": "Biomass productivity",
    "Value": engineering.get(
        "biomass_productivity_g_L_h"
    ),
    "Unit": "g/L/h",
},
                {
    "Metric": "Product productivity",
    "Value": engineering.get(
        "product_productivity_g_L_h"
    ),
    "Unit": "g/L/h",
},
                {
                    "Metric": "Dilution rate",
                    "Value": engineering.get(
                        "dilution_rate_1_h"
                    ),
                    "Unit": "1/h",
                },
                {
                    "Metric": "Volume balance",
                    "Value": engineering.get(
                        "volume_balance_status"
                    ),
                    "Unit": "",
                },
            ]

            st.dataframe(
                analysis_rows,
                width="stretch",
                hide_index=True,
            )

    display_model_trajectories(
        result
    )

    with st.expander("Model details"):
        st.write("Important parameters")
        st.write(
            result.important_parameters
        )

        st.write("Uncertainty")
        st.json(
            result.uncertainty
        )

        st.write("Assumptions")
        st.write(
            result.model_metadata.get(
                "assumptions",
                [],
            )
        )

def _get_parameter_value(state: ProcessState, name: str, default: float | None = None):
    target = name.strip().lower()
    for parameter in state.parameters:
        if parameter.name.strip().lower() == target:
            return parameter.value
    return default


def _upsert_user_parameter(
    state: ProcessState,
    name: str,
    value: float,
    unit: str,
) -> None:
    """Add or replace a process parameter while preserving user provenance."""
    target = name.strip().lower()
    for parameter in state.parameters:
        if parameter.name.strip().lower() == target:
            parameter.value = float(value)
            parameter.unit = unit
            parameter.source = Provenance.USER_PROVIDED
            return

    state.parameters.append(
        Parameter(
            name=name,
            value=float(value),
            unit=unit,
            source=Provenance.USER_PROVIDED,
        )
    )


def display_ecoli_input_form(state: ProcessState, missing_parameters: list[str]) -> bool:
    """Render editable fields for missing E. coli process inputs."""
    st.markdown("### Complete E. coli model inputs")
    st.caption(
        "Enter only the process-specific values that are missing. "
        "These values are recorded as user-provided inputs and are then passed to the model."
    )

    labels = {
        "initial_biomass": ("Initial biomass concentration", "g/L"),
        "initial_substrate": ("Initial substrate concentration", "g/L"),
        "initial_product": ("Initial product concentration", "g/L"),
        "initial_volume": ("Initial reactor volume", "L"),
        "feed_rate": ("Feed rate", "L/h"),
        "feed_substrate": ("Feed substrate concentration", "g/L"),
        "duration": ("Process duration", "h"),
    }

    values = {}
    with st.form("ecoli_missing_inputs_form"):
        columns = st.columns(2)
        for index, parameter_name in enumerate(missing_parameters):
            label, unit = labels.get(
                parameter_name,
                (parameter_name.replace("_", " ").title(), ""),
            )
            default = _get_parameter_value(state, parameter_name, 0.0)
            with columns[index % 2]:
                values[parameter_name] = st.number_input(
                    f"{label} ({unit})" if unit else label,
                    min_value=0.0,
                    value=float(default or 0.0),
                    format="%.6g",
                    key=f"ecoli_input_{parameter_name}",
                )

        submitted = st.form_submit_button(
            "Apply inputs and run E. coli model",
            type="primary",
        )

    if not submitted:
        return False

    for parameter_name, value in values.items():
        _, unit = labels.get(
            parameter_name,
            (parameter_name.replace("_", " ").title(), ""),
        )
        _upsert_user_parameter(
            state,
            parameter_name,
            value,
            unit,
        )

    readiness = get_ecoli_model_readiness(state)
    st.session_state["process_state"] = state
    st.session_state["model_readiness"] = readiness
    st.session_state["model_error"] = None

    if readiness["ready"]:
        try:
            st.session_state["model_result"] = run_from_process_state(state)
        except (TypeError, ValueError) as exc:
            st.session_state["model_result"] = None
            st.session_state["model_error"] = str(exc)
    else:
        st.session_state["model_result"] = None

    return True


def display_scale_up_section(state: ProcessState, result: ModelResult) -> dict | None:
    """Render the first-pass volumetric scale-up screening calculation."""
    st.subheader("Scale-up calculation")
    st.info("**What this means:** Scale converts the model result into an engineering screening basis. It does not certify equipment design. Values generated here remain editable and every scale-up criterion is shown with its assumptions.")
    st.caption(
        "First-pass volumetric scaling from a laboratory process to a target initial working volume. "
        "The current implemented basis preserves concentration and scales volumetric feed rate with volume."
    )

    input_parameters = result.model_metadata.get("input_parameters", {})

    lab_volume = float(
        input_parameters.get(
            "initial_volume",
            _get_parameter_value(state, "initial_volume", 0.0) or 0.0,
        )
    )
    initial_biomass = float(input_parameters.get("initial_biomass", 0.0))
    initial_substrate = float(input_parameters.get("initial_substrate", 0.0))
    initial_product = float(input_parameters.get("initial_product", 0.0))
    feed_rate = float(input_parameters.get("feed_rate", 0.0))
    feed_substrate = float(input_parameters.get("feed_substrate", 0.0))
    duration = float(input_parameters.get("duration", 0.0))

    scale_up_basis = "Constant concentration and volumetric feed rate"

    with st.form("scale_up_form"):
        st.markdown("#### Scale-up basis")
        st.selectbox(
            "Implemented scale-up basis",
            [scale_up_basis],
            index=0,
            help=(
                "This version performs first-pass volumetric scaling. "
                "Constant P/V, tip speed, kLa and mixing-time criteria are not yet calculated."
            ),
        )

        st.markdown("#### Scale-up inputs")
        col1, col2 = st.columns(2)
        with col1:
            scale_lab_volume = st.number_input(
                "Laboratory working volume (L)",
                min_value=0.0001,
                value=lab_volume,
                format="%.6g",
                key="scaleup_lab_volume",
            )
            target_volume = st.number_input(
                "Target initial working volume (L)",
                min_value=0.0001,
                value=lab_volume,
                format="%.6g",
                key="scaleup_target_volume",
                help="Initial working volume at the target scale, before feed addition.",
            )
            max_working_volume = st.number_input(
                "Maximum allowable working volume (L)",
                min_value=0.0,
                value=0.0,
                format="%.6g",
                key="scaleup_max_working_volume",
                help="Enter 0 when vessel capacity is not yet known. The capacity check will then be reported as not evaluated.",
            )
            scale_feed_rate = st.number_input(
                "Laboratory feed rate (L/h)",
                min_value=0.0,
                value=feed_rate,
                format="%.6g",
                disabled=True,
                key="scaleup_feed_rate",
            )
            scale_duration = st.number_input(
                "Process duration (h)",
                min_value=0.0001,
                value=duration,
                format="%.6g",
                disabled=True,
                key="scaleup_duration",
            )
        with col2:
            scale_biomass = st.number_input(
                "Initial biomass concentration (g/L)",
                min_value=0.0,
                value=initial_biomass,
                format="%.6g",
                key="scaleup_biomass",
            )
            scale_substrate = st.number_input(
                "Initial substrate concentration (g/L)",
                min_value=0.0,
                value=initial_substrate,
                format="%.6g",
                key="scaleup_substrate",
            )
            scale_product = st.number_input(
                "Initial product concentration (g/L)",
                min_value=0.0,
                value=initial_product,
                format="%.6g",
                key="scaleup_product",
            )
            scale_feed_substrate = st.number_input(
                "Feed substrate concentration (g/L)",
                min_value=0.0,
                value=feed_substrate,
                format="%.6g",
                key="scaleup_feed_substrate",
            )

        submitted = st.form_submit_button(
            "Calculate scale-up",
            type="primary",
        )

    if not submitted:
        return st.session_state.get("scaleup_result")

    try:
        scaleup_result = calculate_first_pass_scale_up(
            lab_volume_L=scale_lab_volume,
            target_volume_L=target_volume,
            initial_biomass_g_L=scale_biomass,
            initial_substrate_g_L=scale_substrate,
            initial_product_g_L=scale_product,
            feed_rate_L_h=scale_feed_rate,
            feed_substrate_g_L=scale_feed_substrate,
            duration_h=scale_duration,
            max_working_volume_L=max_working_volume,
            scale_up_basis=scale_up_basis,
        )
        st.session_state["scaleup_result"] = scaleup_result
        return scaleup_result
    except ValueError as exc:
        st.error(f"Scale-up calculation could not be completed: {exc}")
        st.session_state["scaleup_result"] = None
        return None


def display_scale_up_result(result: dict | None) -> None:
    if not result:
        return

    st.markdown("#### Scale-up result")
    metrics = st.columns(4)
    metrics[0].metric("Scale factor", f"{result['scale_factor']:.4g}×")
    metrics[1].metric("Scaled feed rate", f"{result['scaled_feed_rate_L_h']:.4g} L/h")
    metrics[2].metric("Scaled feed volume", f"{result['scaled_feed_volume_L']:.4g} L")
    metrics[3].metric("Expected final volume", f"{result['expected_final_volume_L']:.4g} L")

    if result.get("large_scale_factor_flag"):
        st.warning(
            f"Large scale transition flagged: {result['scale_factor']:.4g}×. "
            "The current calculation is a volumetric screening basis and does not establish engineering equivalence at the target scale."
        )

    capacity_status = result.get("vessel_capacity_status", "not_evaluated")
    if capacity_status == "pass":
        st.success(result.get("vessel_capacity_message", "Vessel capacity check passed."))
    elif capacity_status == "check":
        st.error(result.get("vessel_capacity_message", "Vessel capacity requires review."))
    else:
        st.info(result.get("vessel_capacity_message", "Vessel capacity was not evaluated."))

    st.markdown("#### Scaled material balance")
    rows = [
        {"Quantity": "Initial biomass mass", "Value": result["scaled_initial_biomass_mass_g"], "Unit": "g"},
        {"Quantity": "Initial substrate mass", "Value": result["scaled_initial_substrate_mass_g"], "Unit": "g"},
        {"Quantity": "Initial product mass", "Value": result["scaled_initial_product_mass_g"], "Unit": "g"},
        {"Quantity": "Feed volume", "Value": result["scaled_feed_volume_L"], "Unit": "L"},
        {"Quantity": "Feed substrate requirement", "Value": result["scaled_substrate_feed_mass_g"], "Unit": "g"},
        {"Quantity": "Target initial working volume", "Value": result["target_volume_L"], "Unit": "L"},
        {"Quantity": "Expected final volume", "Value": result["expected_final_volume_L"], "Unit": "L"},
    ]
    st.dataframe(rows, width="stretch", hide_index=True)


def display_engineering_scale_up_section(result: dict | None, model_result: ModelResult | None = None) -> dict | None:
    """Render an automated-but-editable Phase 2 engineering scale-up workflow."""
    st.markdown("---")
    st.subheader("Phase 2 — Engineering scale-up calculations")
    st.caption(
        "ScaleUp AI generates a screening basis from the process scale and any laboratory data supplied. "
        "Calculated target values remain editable; assumptions are shown explicitly and are never presented as measured equipment data."
    )

    if result is None:
        st.info("Run the first-pass scale-up before opening Phase 2 engineering calculations.")
        return None

    lab_volume = float(result["lab_volume_L"])
    target_volume = float(result["target_volume_L"])

    # ---------------------------------------------------------------
    # Engineering basis controls
    # ---------------------------------------------------------------
    b1, b2, b3 = st.columns(3)
    with b1:
        geometry_basis = st.selectbox(
            "Geometry basis",
            ["Geometric similarity (cylindrical screening)"],
            key="eng_geometry_basis",
        )
    with b2:
        geometry_ht = st.number_input(
            "H/T ratio — liquid height / vessel diameter",
            min_value=0.1,
            value=float(st.session_state.get("eng_geometry_ht", 2.0)),
            step=0.1,
            format="%.3f",
            key="eng_geometry_ht",
            help="Liquid height / vessel diameter used only when measured geometry is not supplied.",
        )
    with b3:
        geometry_dt = st.number_input(
            "D/T ratio — impeller diameter / vessel diameter",
            min_value=0.05,
            value=float(st.session_state.get("eng_geometry_dt", 1.0 / 3.0)),
            step=0.01,
            format="%.3f",
            key="eng_geometry_dt",
            help="Impeller diameter / vessel diameter used only when measured impeller geometry is not supplied.",
        )

    with st.expander("What do H/T and D/T mean?", expanded=True):
        st.markdown(
            "**H/T and D/T are dimensionless reactor-geometry ratios used to estimate missing equipment dimensions.**"
        )
        ht1, ht2 = st.columns(2)
        with ht1:
            st.markdown(
                "**H/T — liquid height / tank diameter**\n\n"
                "- **H** = liquid height in the vessel\n"
                "- **T** = tank/vessel diameter\n"
                "- Example: H/T = 2 means the liquid height is twice the tank diameter.\n\n"
                "For a cylindrical screening geometry, ScaleUp uses this ratio with the working volume to estimate vessel diameter and liquid height."
            )
        with ht2:
            st.markdown(
                "**D/T — impeller diameter / tank diameter**\n\n"
                "- **D** = impeller diameter\n"
                "- **T** = tank/vessel diameter\n"
                "- Example: D/T = 0.333 means the impeller diameter is about one-third of the tank diameter.\n\n"
                "ScaleUp uses this ratio to estimate impeller diameter when measured impeller geometry is not supplied."
            )
        st.info(
            "**Why these ratios matter:** they provide a declared geometric-similarity basis for the first-pass scale-up. "
            "They are engineering assumptions, not universal design rules. If actual vessel or impeller dimensions are known, enter those values and override the estimates. "
            "Changing H/T or D/T changes the automatically calculated geometry and any downstream mixing calculations that depend on it."
        )
        st.caption(
            "Current screening defaults: H/T = 2.0 and D/T = 0.333. Both ratios are dimensionless."
        )

    m1, m2 = st.columns(2)
    with m1:
        mixing_basis = st.selectbox(
            "Mixing scale-up basis",
            ["constant_tip_speed", "constant_P_over_V"],
            format_func=lambda x: "Constant tip speed" if x == "constant_tip_speed" else "Constant P/V",
            key="eng_mixing_basis",
        )
    with m2:
        aeration_basis = st.selectbox(
            "Aeration scale-up basis",
            ["constant_vvm"],
            format_func=lambda x: "Constant vvm",
            key="eng_aeration_basis",
        )

    # Generate the initial engineering basis automatically once. The explicit
    # regenerate button below can be used after laboratory inputs are changed.
    if not st.session_state.get("engineering_scaleup_basis_initialized", False):
        try:
            initial_suggestion = suggest_engineering_inputs(
                lab_volume_L=lab_volume, target_volume_L=target_volume,
                geometry_H_over_T=geometry_ht, geometry_D_over_T=geometry_dt,
                mixing_basis=mixing_basis, aeration_basis=aeration_basis,
            )
            geo = initial_suggestion["geometry"]
            for key, value in {
                "eng_lab_vessel_d": geo["lab_vessel_diameter_m"],
                "eng_lab_height": geo["lab_liquid_height_m"],
                "eng_lab_impeller_d": geo["lab_impeller_diameter_m"],
                "eng_target_vessel_d": geo["target_vessel_diameter_m"],
                "eng_target_height": geo["target_liquid_height_m"],
                "eng_target_impeller_d": geo["target_impeller_diameter_m"],
            }.items():
                st.session_state[key] = float(value)
            st.session_state["engineering_scaleup_suggestion"] = initial_suggestion
            st.session_state["engineering_scaleup_basis_initialized"] = True
        except ValueError:
            pass

    st.markdown("#### Laboratory data / automatically estimated basis")
    st.caption(
        "Enter measured laboratory values where available. If geometry is left blank, ScaleUp AI estimates it from working volume and the stated H/T and D/T assumptions."
    )
    l1, l2 = st.columns(2)
    with l1:
        lab_vessel_d = st.number_input(
            "Lab vessel diameter (m)", min_value=0.0,
            value=float(st.session_state.get("eng_lab_vessel_d", 0.0)),
            step=0.001, format="%.6g", key="eng_lab_vessel_d",
        )
        lab_height = st.number_input(
            "Lab liquid height (m)", min_value=0.0,
            value=float(st.session_state.get("eng_lab_height", 0.0)),
            step=0.001, format="%.6g", key="eng_lab_height",
        )
        lab_impeller_d = st.number_input(
            "Lab impeller diameter (m)", min_value=0.0,
            value=float(st.session_state.get("eng_lab_impeller_d", 0.0)),
            step=0.001, format="%.6g", key="eng_lab_impeller_d",
        )
    with l2:
        lab_rpm = st.number_input(
            "Lab agitation (rpm)", min_value=0.0,
            value=float(st.session_state.get("eng_lab_rpm", 0.0)),
            step=10.0, format="%.6g", key="eng_lab_rpm",
            help="Enter the measured laboratory agitation. Target agitation is then calculated from the selected mixing basis.",
        )
        lab_np = st.number_input(
            "Lab power number, Np (optional)", min_value=0.0,
            value=float(st.session_state.get("eng_lab_np", 0.0)),
            step=0.1, format="%.6g", key="eng_lab_np",
            help="Do not guess this value; it depends on impeller type and flow regime.",
        )
        density = st.number_input(
            "Fluid density (kg/m³)", min_value=0.001,
            value=float(st.session_state.get("eng_density", 1000.0)),
            step=10.0, format="%.6g", key="eng_density",
        )
        lab_viscosity = st.number_input(
            "Lab dynamic viscosity (Pa·s, optional)", min_value=0.0,
            value=float(st.session_state.get("eng_lab_viscosity", 0.0)),
            step=0.0001, format="%.6g", key="eng_lab_viscosity",
            help="Used to calculate impeller Reynolds number. Enter a measured or justified process value; do not guess."
        )
        target_viscosity = st.number_input(
            "Target dynamic viscosity (Pa·s, optional)", min_value=0.0,
            value=float(st.session_state.get("eng_target_viscosity", 0.0)),
            step=0.0001, format="%.6g", key="eng_target_viscosity",
            help="Used for target Reynolds number and constant-Re comparison when it differs from the laboratory fluid."
        )
        lab_mixing_time = st.number_input(
            "Lab mixing time (s, optional)", min_value=0.0,
            value=float(st.session_state.get("eng_lab_mixing_time", 0.0)),
            step=1.0, format="%.6g", key="eng_lab_mixing_time",
            help="Measured mixing time can be used for a dimensionless mixing-time screening estimate."
        )
        lab_gas = st.number_input(
            "Lab gas flow (L/min)", min_value=0.0,
            value=float(st.session_state.get("eng_lab_gas", 0.0)),
            step=0.1, format="%.6g", key="eng_lab_gas",
            help="Required to automatically scale target gas flow at constant vvm.",
        )

    # Propagate target agitation/gas flow automatically from current laboratory
    # operating inputs. Manual target overrides are still possible after the
    # calculated basis has been generated.
    try:
        live_suggestion = suggest_engineering_inputs(
            lab_volume_L=lab_volume,
            target_volume_L=target_volume,
            lab_vessel_diameter_m=lab_vessel_d or None,
            lab_liquid_height_m=lab_height or None,
            lab_impeller_diameter_m=lab_impeller_d or None,
            lab_agitation_rpm=lab_rpm or None,
            geometry_H_over_T=geometry_ht,
            geometry_D_over_T=geometry_dt,
            mixing_basis=mixing_basis,
            lab_gas_flow_L_min=lab_gas or None,
            aeration_basis=aeration_basis,
        )
        if not st.session_state.get("eng_target_rpm_manual", False):
            st.session_state["eng_target_rpm"] = float(live_suggestion["mixing"]["target_agitation_rpm"] or 0.0)
        if not st.session_state.get("eng_target_gas_manual", False):
            st.session_state["eng_target_gas"] = float(live_suggestion["aeration"]["target_gas_flow_L_min"] or 0.0)
    except ValueError:
        live_suggestion = None

    # Use a button callback for regeneration. Streamlit executes the callback
    # before the next script run, so widget-backed session-state keys can be
    # updated safely here. Mutating a widget key later in the same run would
    # raise StreamlitWidgetAlreadyInstantiatedError.
    def _regenerate_engineering_basis() -> None:
        try:
            suggested = suggest_engineering_inputs(
                lab_volume_L=lab_volume,
                target_volume_L=target_volume,
                lab_vessel_diameter_m=float(st.session_state.get("eng_lab_vessel_d", 0.0)) or None,
                lab_liquid_height_m=float(st.session_state.get("eng_lab_height", 0.0)) or None,
                lab_impeller_diameter_m=float(st.session_state.get("eng_lab_impeller_d", 0.0)) or None,
                lab_agitation_rpm=float(st.session_state.get("eng_lab_rpm", 0.0)) or None,
                geometry_H_over_T=float(st.session_state.get("eng_geometry_ht", geometry_ht)),
                geometry_D_over_T=float(st.session_state.get("eng_geometry_dt", geometry_dt)),
                mixing_basis=st.session_state.get("eng_mixing_basis", mixing_basis),
                lab_gas_flow_L_min=float(st.session_state.get("eng_lab_gas", 0.0)) or None,
                aeration_basis=st.session_state.get("eng_aeration_basis", aeration_basis),
            )
            geo = suggested["geometry"]
            mix = suggested["mixing"]
            aer = suggested["aeration"]
            for key, value in {
                "eng_lab_vessel_d": geo["lab_vessel_diameter_m"],
                "eng_lab_height": geo["lab_liquid_height_m"],
                "eng_lab_impeller_d": geo["lab_impeller_diameter_m"],
                "eng_target_vessel_d": geo["target_vessel_diameter_m"],
                "eng_target_height": geo["target_liquid_height_m"],
                "eng_target_impeller_d": geo["target_impeller_diameter_m"],
                "eng_target_rpm": mix["target_agitation_rpm"] or 0.0,
                "eng_target_gas": aer["target_gas_flow_L_min"] or 0.0,
            }.items():
                st.session_state[key] = float(value)
            st.session_state["engineering_scaleup_suggestion"] = suggested
            st.session_state["engineering_scaleup_basis_initialized"] = True
            st.session_state["eng_target_rpm_manual"] = False
            st.session_state["eng_target_gas_manual"] = False
            st.session_state["engineering_basis_regenerated"] = True
            st.session_state.pop("engineering_basis_regeneration_error", None)
        except ValueError as exc:
            st.session_state["engineering_basis_regeneration_error"] = str(exc)

    st.button(
        "Generate / update engineering basis",
        type="primary",
        key="generate_engineering_basis",
        on_click=_regenerate_engineering_basis,
    )
    if st.session_state.pop("engineering_basis_regenerated", False):
        st.success("Engineering basis generated. The calculated target values below remain editable.")
    if error := st.session_state.pop("engineering_basis_regeneration_error", None):
        st.error(f"Engineering basis could not be generated: {error}")

    suggestion = st.session_state.get("engineering_scaleup_suggestion")
    if suggestion:
        st.markdown("#### Automatically generated engineering basis")
        st.caption(
            "These values are calculated from working volume plus the stated scale-up assumptions. Change any target value manually if you have actual equipment data."
        )
        s1, s2 = st.columns(2)
        with s1:
            st.write("**Geometry calculation**")
            st.code(
                "V = π/4 × T² × H\n"
                f"H/T = H ÷ T = {geometry_ht:g}\n"
                f"D/T = D_impeller ÷ T = {geometry_dt:g}\n"
                f"Lab volume = {lab_volume:g} L\n"
                f"Target volume = {target_volume:g} L",
                language="text",
            )
        with s2:
            st.write("**Scale-up basis**")
            st.code(
                f"Mixing: {'constant tip speed' if mixing_basis == 'constant_tip_speed' else 'constant P/V'}\n"
                f"Aeration: constant vvm\n"
                f"Scale factor = {target_volume / lab_volume:g}×",
                language="text",
            )

    # ---------------------------------------------------------------
    # Editable target values + optional oxygen/heat data
    # ---------------------------------------------------------------
    st.markdown("#### Engineering-data provenance")
    st.caption(
        "Geometry is automatically calculated from working volume when measured dimensions are unavailable. "
        "Agitation and gas flow remain dependent on measured laboratory operating data. Target values may be manually overridden."
    )
    provenance_preview = [
        {"Item": "Laboratory geometry", "Source": "Calculated from working volume + H/T + D/T"},
        {"Item": "Target geometry", "Source": "Calculated from target working volume + H/T + D/T"},
        {"Item": "Laboratory agitation", "Source": "User/equipment data required"},
        {"Item": "Target agitation", "Source": "Calculated from laboratory agitation + selected basis"},
        {"Item": "Laboratory gas flow", "Source": "User/equipment data required"},
        {"Item": "Target gas flow", "Source": "Calculated from laboratory gas flow + constant-vvm basis"},
        {"Item": "kLa / Np / heat-transfer data", "Source": "Measured/supplied equipment/process data"},
    ]
    st.dataframe(provenance_preview, width="stretch", hide_index=True)

    st.markdown("#### Editable target engineering values")
    t1, t2 = st.columns(2)
    with t1:
        target_vessel_d = st.number_input(
            "Target vessel diameter (m)", min_value=0.0,
            value=float(st.session_state.get("eng_target_vessel_d", 0.0)),
            step=0.001, format="%.6g", key="eng_target_vessel_d",
        )
        target_height = st.number_input(
            "Target liquid height (m)", min_value=0.0,
            value=float(st.session_state.get("eng_target_height", 0.0)),
            step=0.001, format="%.6g", key="eng_target_height",
        )
        target_impeller_d = st.number_input(
            "Target impeller diameter (m)", min_value=0.0,
            value=float(st.session_state.get("eng_target_impeller_d", 0.0)),
            step=0.001, format="%.6g", key="eng_target_impeller_d",
        )
        target_rpm = st.number_input(
            "Target agitation (rpm)", min_value=0.0,
            value=float(st.session_state.get("eng_target_rpm", 0.0)),
            step=10.0, format="%.6g", key="eng_target_rpm",
            on_change=lambda: st.session_state.__setitem__("eng_target_rpm_manual", True),
        )
        target_np = st.number_input(
            "Target power number, Np (optional)", min_value=0.0,
            value=float(st.session_state.get("eng_target_np", 0.0)),
            step=0.1, format="%.6g", key="eng_target_np",
        )
    with t2:
        target_gas = st.number_input(
            "Target gas flow (L/min)", min_value=0.0,
            value=float(st.session_state.get("eng_target_gas", 0.0)),
            step=0.1, format="%.6g", key="eng_target_gas",
            on_change=lambda: st.session_state.__setitem__("eng_target_gas_manual", True),
        )
        lab_kla = st.number_input(
            "Lab kLa (1/h, optional)", min_value=0.0,
            value=float(st.session_state.get("eng_lab_kla", 0.0)),
            step=1.0, format="%.6g", key="eng_lab_kla",
        )
        target_kla = st.number_input(
            "Target kLa (1/h, optional)", min_value=0.0,
            value=float(st.session_state.get("eng_target_kla", 0.0)),
            step=1.0, format="%.6g", key="eng_target_kla",
        )
        cstar = st.number_input(
            "Target oxygen saturation C* (g/L, optional)", min_value=0.0,
            value=float(st.session_state.get("eng_cstar", 0.0)),
            step=0.0001, format="%.6g", key="eng_cstar",
        )
        do = st.number_input(
            "Target dissolved oxygen CL (g/L, optional)", min_value=0.0,
            value=float(st.session_state.get("eng_do", 0.0)),
            step=0.0001, format="%.6g", key="eng_do",
        )
        qo2 = st.number_input(
            "Specific oxygen uptake qO2 (g O2/g biomass/h, optional)",
            min_value=0.0,
            value=float(st.session_state.get("eng_qo2", 0.0)),
            step=0.01, format="%.6g", key="eng_qo2",
            help="Process-specific oxygen uptake rate. Enter a measured or justified value; ScaleUp AI does not infer qO2 from the Monod model."
        )

    st.info(
        "OUR is now connected to the scientific biomass trajectory when qO2 is supplied: "
        "OUR = qO2 × X. The target check assumes concentration-preserving scale-up, "
        "so target volumetric OUR follows the simulated biomass concentration. qO2 must be measured or justified for the process; it is not a universal E. coli constant."
    )

    st.markdown("#### Heat-generation and heat-transfer data (optional)")
    st.caption(
        "If qO2 and a process-specific heat-release basis are supplied, ScaleUp AI estimates peak metabolic heat from the same oxygen-demand calculation used above. "
        "A manual total heat duty can override the calculated duty for the capacity comparison; the calculated metabolic heat remains visible for traceability."
    )
    h1, h2 = st.columns(2)
    with h1:
        heat_release = st.number_input(
            "Heat release per oxygen consumed (kJ/g O2, optional)",
            min_value=0.0,
            value=float(st.session_state.get("eng_heat_release", 0.0)),
            step=10.0, format="%.6g", key="eng_heat_release",
            help="Process-specific or literature-supported heat-release basis. ScaleUp AI does not assume a universal value.",
        )
        additional_heat = st.number_input(
            "Additional target heat load (W, optional)",
            min_value=0.0,
            value=float(st.session_state.get("eng_additional_heat", 0.0)),
            step=100.0, format="%.6g", key="eng_additional_heat",
            help="Non-metabolic heat load to add to calculated metabolic heat when no manual total duty is supplied.",
        )
        heat_duty = st.number_input(
            "Manual total target heat duty (W, optional)",
            min_value=0.0,
            value=float(st.session_state.get("eng_heat_duty", 0.0)),
            step=100.0, format="%.6g", key="eng_heat_duty",
            help="Optional total heat-removal duty. When supplied, it is used for the capacity comparison instead of the calculated metabolic + additional-load duty.",
        )
    with h2:
        heat_u = st.number_input("Overall heat-transfer coefficient U (W/m²·K)", min_value=0.0, value=float(st.session_state.get("eng_heat_u", 0.0)), step=10.0, format="%.6g", key="eng_heat_u")
        heat_area = st.number_input("Available heat-transfer area (m²)", min_value=0.0, value=float(st.session_state.get("eng_heat_area", 0.0)), step=1.0, format="%.6g", key="eng_heat_area")
        delta_t = st.number_input("Available temperature driving force ΔT (K)", min_value=0.0, value=float(st.session_state.get("eng_delta_t", 0.0)), step=1.0, format="%.6g", key="eng_delta_t")

    st.caption("Phase 2 recalculates automatically from the current inputs. Target values generated from the selected basis can be manually overridden.")
    try:
        def optional(value: float) -> float | None:
            return value if value > 0 else None

        engineering_result = calculate_engineering_scale_up(
            lab_volume_L=lab_volume,
            target_volume_L=target_volume,
            lab_vessel_diameter_m=optional(lab_vessel_d),
            target_vessel_diameter_m=optional(target_vessel_d),
            lab_liquid_height_m=optional(lab_height),
            target_liquid_height_m=optional(target_height),
            lab_impeller_diameter_m=optional(lab_impeller_d),
            target_impeller_diameter_m=optional(target_impeller_d),
            lab_agitation_rpm=optional(lab_rpm),
            target_agitation_rpm=optional(target_rpm),
            lab_power_number=optional(lab_np),
            target_power_number=optional(target_np),
            fluid_density_kg_m3=density,
            lab_viscosity_Pa_s=optional(lab_viscosity),
            target_viscosity_Pa_s=optional(target_viscosity),
            lab_mixing_time_s=optional(lab_mixing_time),
            lab_gas_flow_L_min=optional(lab_gas),
            target_gas_flow_L_min=optional(target_gas),
            lab_kla_1_h=optional(lab_kla),
            target_kla_1_h=optional(target_kla),
            target_cstar_g_L=optional(cstar),
            target_do_g_L=optional(do),
            target_heat_duty_W=optional(heat_duty),
            target_overall_heat_transfer_U_W_m2K=optional(heat_u),
            target_heat_transfer_area_m2=optional(heat_area),
            target_delta_T_K=optional(delta_t),
            heat_release_kJ_per_gO2=optional(heat_release),
            additional_heat_load_W=optional(additional_heat),
            qO2_g_gX_h=optional(qo2),
            biomass_trajectory_g_L=(
                (model_result.predictions.get("biomass_g_L", []) if model_result is not None else [])
                or None
            ),
            trajectory_time_h=(
                (model_result.predictions.get("time_h", []) if model_result is not None else [])
                or None
            ),
        )
        engineering_result["auto_basis"] = suggestion or live_suggestion
        st.session_state["engineering_scaleup_result"] = engineering_result
    except ValueError as exc:
        st.error(f"Phase 2 calculation could not be completed: {exc}")
        st.session_state["engineering_scaleup_result"] = None

    engineering_result = st.session_state.get("engineering_scaleup_result")
    if not engineering_result:
        return None

    st.markdown("#### Phase 2 calculation results")
    st.caption("The table below exposes the calculations used by the engineering screening rather than hiding them behind a pass/fail label.")

    geometry = engineering_result["geometry"]
    metric_specs = [
        ("Scale factor", target_volume / lab_volume, "×"),
        ("Vessel diameter ratio", geometry["vessel_diameter_ratio"], "×"),
        ("Liquid height ratio", geometry["liquid_height_ratio"], "×"),
        ("Impeller diameter ratio", geometry["impeller_diameter_ratio"], "×"),
        ("Lab H/T", geometry["H_over_T_lab"], "×"),
        ("Target H/T", geometry["H_over_T_target"], "×"),
        ("Lab D/T", geometry["D_over_T_lab"], "×"),
        ("Target D/T", geometry["D_over_T_target"], "×"),
        ("Lab tip speed", engineering_result["tip_speed"]["lab_m_s"], "m/s"),
        ("Target tip speed", engineering_result["tip_speed"]["target_m_s"], "m/s"),
        ("Tip-speed ratio", engineering_result["tip_speed"]["ratio"], "×"),
        ("Lab power", engineering_result["power"]["lab_W"], "W"),
        ("Target power", engineering_result["power"]["target_W"], "W"),
        ("Lab P/V", engineering_result["power"]["lab_W_m3"], "W/m³"),
        ("Target P/V", engineering_result["power"]["target_W_m3"], "W/m³"),
        ("P/V ratio", engineering_result["power"]["ratio"], "×"),
        ("Lab vvm", engineering_result["aeration"]["lab_vvm"], "vvm"),
        ("Target vvm", engineering_result["aeration"]["target_vvm"], "vvm"),
        ("vvm ratio", engineering_result["aeration"]["ratio"], "×"),
        ("Lab impeller Reynolds number", engineering_result["similarity"]["lab_reynolds"], "—"),
        ("Target impeller Reynolds number", engineering_result["similarity"]["target_reynolds"], "—"),
        ("Reynolds-number ratio", engineering_result["similarity"]["reynolds_ratio"], "×"),
        ("Lab Froude number", engineering_result["similarity"]["lab_froude"], "—"),
        ("Target Froude number", engineering_result["similarity"]["target_froude"], "—"),
        ("Froude-number ratio", engineering_result["similarity"]["froude_ratio"], "×"),
        ("Predicted target mixing time", engineering_result["similarity"]["predicted_target_mixing_time_s"], "s"),
        ("kLa ratio", engineering_result["oxygen"]["kla_ratio"], "×"),
        ("Target OTR", engineering_result["oxygen"]["target_otr_g_L_h"], "g/L/h"),
        ("Peak lab OUR", engineering_result["oxygen_demand"]["lab_peak_our_g_L_h"], "g/L/h"),
        ("Peak target OUR", engineering_result["oxygen_demand"]["target_peak_our_g_L_h"], "g/L/h"),
        ("Peak target OUR (total)", engineering_result["oxygen_demand"]["target_peak_our_g_h"], "g/h"),
        ("OTR / peak OUR", engineering_result["oxygen_demand"]["transfer_capacity_ratio"], "×"),
        ("Critical biomass concentration", engineering_result["oxygen_demand"]["target_critical_biomass_g_L"], "g/L"),
        ("Target oxygen-transfer margin", engineering_result["oxygen_demand"]["transfer_margin_g_L_h"], "g/L/h"),
        ("Peak metabolic heat generation", engineering_result["heat"]["target_metabolic_heat_generation_W"], "W"),
        ("Additional heat load", engineering_result["heat"]["additional_heat_load_W"], "W"),
        ("Target heat duty used for capacity check", engineering_result["heat"]["target_heat_duty_basis_W"], "W"),
        ("Target heat-removal capacity", engineering_result["heat"]["target_heat_capacity_W"], "W"),
        ("Heat capacity / duty", engineering_result["heat"]["capacity_to_duty_ratio"], "×"),
        ("Heat-removal margin", engineering_result["heat"]["heat_margin_W"], "W"),
        ("Thermal critical biomass concentration", engineering_result["heat"]["thermal_critical_biomass_g_L"], "g/L"),
    ]
    metric_rows = [
        {"Metric": label, "Value": "Not evaluated" if value is None else f"{value:.6g}", "Unit": unit}
        for label, value, unit in metric_specs
    ]
    st.dataframe(metric_rows, width="stretch", hide_index=True)

    st.markdown("#### Classical mixing scale-up comparison")
    st.caption(
        "For geometrically similar stirred tanks, classical similarity analysis compares alternative criteria rather than assuming one criterion is universally correct. "
        "The table below shows the target agitation each criterion would imply from the current laboratory basis."
    )
    criteria = engineering_result["similarity"]["criteria_comparison_target_rpm"]
    criterion_rows = [
        {
            "Criterion": "Constant tip speed",
            "Target agitation (rpm)": criteria.get("constant_tip_speed"),
            "What it preserves": "Impeller tip velocity",
        },
        {
            "Criterion": "Constant P/V (turbulent screening)",
            "Target agitation (rpm)": criteria.get("constant_P_over_V"),
            "What it preserves": "Agitator power per unit volume",
        },
        {
            "Criterion": "Constant impeller Reynolds number",
            "Target agitation (rpm)": criteria.get("constant_Re"),
            "What it preserves": "Reynolds-number similarity",
        },
        {
            "Criterion": "Constant Froude number",
            "Target agitation (rpm)": criteria.get("constant_Fr"),
            "What it preserves": "Free-surface/gravity similarity",
        },
    ]
    st.dataframe(criterion_rows, width="stretch", hide_index=True)
    st.info(
        "These criteria can give different target speeds. ScaleUp AI therefore treats the selected basis as a process decision, "
        "not as a universal rule. Constant P/V is shown as a turbulent, geometrically similar screening relation and should not be used blindly outside its assumptions."
    )

    st.markdown("#### Oxygen demand versus transfer capacity")
    oxygen_demand = engineering_result["oxygen_demand"]
    oxygen_rows = [
        {"Quantity": "Specific oxygen uptake qO2", "Value": oxygen_demand["qO2_g_gX_h"], "Unit": "g O2/g biomass/h"},
        {"Quantity": "Peak biomass concentration", "Value": oxygen_demand["lab_peak_biomass_g_L"], "Unit": "g/L"},
        {"Quantity": "Peak OUR", "Value": oxygen_demand["target_peak_our_g_L_h"], "Unit": "g O2/L/h"},
        {"Quantity": "Target OTR", "Value": engineering_result["oxygen"]["target_otr_g_L_h"], "Unit": "g O2/L/h"},
        {"Quantity": "OTR / OUR capacity ratio", "Value": oxygen_demand["transfer_capacity_ratio"], "Unit": "×"},
        {"Quantity": "OTR − OUR margin", "Value": oxygen_demand["transfer_margin_g_L_h"], "Unit": "g O2/L/h"},
        {"Quantity": "Critical biomass concentration", "Value": oxygen_demand["target_critical_biomass_g_L"], "Unit": "g/L"},
        {"Quantity": "Screening status", "Value": oxygen_demand["status"], "Unit": ""},
    ]
    st.dataframe(oxygen_rows, width="stretch", hide_index=True)
    if oxygen_demand["status"] == "screening_capacity_available":
        st.success("Screening oxygen-transfer capacity is above the calculated peak OUR under the stated inputs.")
    elif oxygen_demand["status"] == "screening_capacity_exceeded":
        st.warning("Calculated peak OUR exceeds the supplied target OTR on this screening basis; oxygen-transfer limitation requires engineering review.")
    elif oxygen_demand["status"] == "our_calculated_otr_not_available":
        st.info("OUR was calculated, but target OTR is not available because kLa, C* and CL are incomplete.")
    else:
        st.info("OUR is not evaluated until a process-specific qO2 and scientific biomass trajectory are available.")

    st.markdown("#### Thermal generation versus removal capacity")
    heat = engineering_result["heat"]
    heat_rows = [
        {"Quantity": "Heat release basis", "Value": heat["heat_release_kJ_per_gO2"], "Unit": "kJ/g O2"},
        {"Quantity": "Peak metabolic heat generation", "Value": heat["target_metabolic_heat_generation_W"], "Unit": "W"},
        {"Quantity": "Additional heat load", "Value": heat["additional_heat_load_W"], "Unit": "W"},
        {"Quantity": "Calculated total heat generation", "Value": heat["target_total_heat_generation_W"], "Unit": "W"},
        {"Quantity": "Heat duty used for capacity check", "Value": heat["target_heat_duty_basis_W"], "Unit": "W"},
        {"Quantity": "Heat-removal capacity", "Value": heat["target_heat_capacity_W"], "Unit": "W"},
        {"Quantity": "Capacity / duty ratio", "Value": heat["capacity_to_duty_ratio"], "Unit": "×"},
        {"Quantity": "Heat-removal margin", "Value": heat["heat_margin_W"], "Unit": "W"},
        {"Quantity": "Thermal critical biomass concentration", "Value": heat["thermal_critical_biomass_g_L"], "Unit": "g/L"},
        {"Quantity": "Generation basis status", "Value": heat["generation_status"], "Unit": ""},
        {"Quantity": "Screening status", "Value": heat["status"], "Unit": ""},
    ]
    st.dataframe(heat_rows, width="stretch", hide_index=True)
    if heat["status"] == "capacity_available":
        st.success("The stated/calculated target heat-removal duty is within the supplied U·A·ΔT screening capacity.")
    elif heat["status"] == "capacity_check":
        st.warning("The target heat-removal duty exceeds the supplied U·A·ΔT screening capacity; thermal design requires engineering review.")
    elif heat["status"] == "duty_calculated_capacity_not_available":
        st.info("A target heat duty was calculated, but heat-removal capacity is not available because U, area and ΔT are incomplete.")
    else:
        st.info("Thermal generation/removal is not fully evaluated until a heat-release basis or heat duty and the required heat-transfer inputs are supplied.")

    st.markdown("#### Similarity diagnostics")
    st.markdown(
        "**Reynolds number:** Reᵢ = ρND²/μ. It indicates the mixing flow regime and is useful for checking whether the lab and target systems are in comparable regimes.  "
        "**Froude number:** Fr = N²D/g. It becomes particularly relevant when free-surface vortexing/gravity effects matter, especially in unbaffled systems."
    )

    st.markdown("#### Calculation basis and provenance")
    provenance_rows = [
        {"Item": "Geometry", "Basis": engineering_result.get("basis", ""), "Source": "Calculated / user supplied"},
        {"Item": "Geometry assumptions", "Basis": f"H/T = {geometry_ht:g}; D/T = {geometry_dt:g}; cylindrical screening geometry", "Source": "User-selected assumption"},
        {"Item": "Mixing", "Basis": mixing_basis, "Source": "User-selected criterion"},
        {"Item": "Aeration", "Basis": aeration_basis, "Source": "User-selected criterion"},
    ]
    st.dataframe(provenance_rows, width="stretch", hide_index=True)

    st.markdown("#### Engineering screening status")
    status_labels = {
        "reactor_geometry": "Reactor geometry",
        "tip_speed": "Tip speed",
        "power_per_volume": "Power per volume",
        "aeration_vvm": "Aeration (vvm)",
        "kla": "kLa",
        "oxygen_transfer": "Oxygen transfer / OTR",
        "heat_transfer": "Heat transfer",
    }
    status_rows = [
        {"Criterion": label, "Status": engineering_result["checks"].get(key, "not_evaluated")}
        for key, label in status_labels.items()
    ]
    st.dataframe(status_rows, width="stretch", hide_index=True)

    with st.expander("Show equations used"):
        st.code(
            "Geometry: V = π/4 × T² × H  (T = tank diameter)\n"
            "H/T: liquid height H ÷ tank diameter T\n"
            "D/T: impeller diameter D_impeller ÷ tank diameter T\n"
            "Tip speed: u_tip = π × D × N\n"
            "Power: P = Np × ρ × N³ × D⁵\n"
            "Power density: P/V\n"
            "Aeration: vvm = gas flow / working volume\n"
            "Impeller Reynolds number: Reᵢ = ρND²/μ\n"
            "Froude number: Fr = N²D/g\n"
            "Classical constant-tip-speed comparison: N₂/N₁ = D₁/D₂\n"
            "Classical constant-P/V comparison (geometric + turbulent screening): N₂/N₁ = (D₁/D₂)^(2/3)\n"
            "Constant-Re comparison: N₂/N₁ = (D₁/D₂)² × (μ₂/μ₁) × (ρ₁/ρ₂)\n"
            "Constant-Fr comparison: N₂/N₁ = sqrt(D₁/D₂)\n"
            "Oxygen transfer: OTR = kLa × (C* − CL)\n"
            "Oxygen uptake: OUR = qO2 × X\n"
            "Oxygen-transfer margin = OTR − OUR\n"
            "Critical biomass concentration: Xcrit = OTR/qO2\n"
            "Metabolic heat generation: Q_met = OUR_total × heat_release / 3.6\n"
            "Total calculated heat duty: Q_duty = Q_met + Q_additional\n"
            "Heat removal capacity: Q = U × A × ΔT\n"
            "Thermal critical biomass: Xcrit,thermal = (Q_capacity − Q_additional) / (qO2 × heat_release × V / 3.6)",
            language="text",
        )

    for warning in engineering_result["warnings"]:
        st.info(warning)

    return engineering_result



def display_phase3_assessment(
    scaleup_result: dict | None,
    engineering_result: dict | None,
    model_result: ModelResult | None,
) -> dict | None:
    """Render Phase 3 integrated engineering assessment."""
    if scaleup_result is None or engineering_result is None:
        return None

    st.markdown("---")
    st.subheader("Phase 3 — Integrated engineering scale-up assessment")
    st.info("**What this means:** Phase 3 combines the available material balance, mixing, oxygen-transfer and thermal checks. A check marked as unresolved means the necessary engineering data were not available; it is not a pass or fail.")
    st.caption(
        "Phase 3 integrates the Phase 1 material/volume basis, Phase 2 engineering calculations, "
        "and available scientific-model outputs. It is an engineering screening assessment, "
        "not final equipment design or process validation."
    )

    st.markdown("#### Assessment basis")
    b1, b2, b3 = st.columns(3)
    with b1:
        st.metric("Scale factor", f"{scaleup_result.get('scale_factor', 0):.4g}×")
    with b2:
        st.metric("Initial target volume", f"{scaleup_result.get('target_volume_L', 0):.6g} L")
    with b3:
        st.metric("Expected final volume", f"{scaleup_result.get('expected_final_volume_L', 0):.6g} L")

    a1, a2 = st.columns(2)
    with a1:
        st.write(f"**Material/volume basis:** {scaleup_result.get('scale_up_basis', 'Not specified')}")
    with a2:
        st.write(f"**Mixing basis:** {st.session_state.get('eng_mixing_basis', 'Not specified')}")

    st.markdown("#### Optional target constraints")
    c1, c2 = st.columns(2)
    with c1:
        max_target_feed = st.number_input(
            "Maximum target feed rate (L/h)",
            min_value=0.0,
            value=float(st.session_state.get("phase3_max_feed", 0.0)),
            step=0.1,
            format="%.6g",
            key="phase3_max_feed",
            help="Leave at 0 if no maximum target feed-rate constraint is available.",
        )
    with c2:
        max_biomass = st.number_input(
            "Maximum biomass concentration (g/L)",
            min_value=0.0,
            value=float(st.session_state.get("phase3_max_biomass", 0.0)),
            step=0.1,
            format="%.6g",
            key="phase3_max_biomass",
            help="Leave at 0 if no maximum biomass concentration constraint is available.",
        )

    try:
        assessment = assess_integrated_scale_up(
            scaleup_result=scaleup_result,
            engineering_result=engineering_result,
            model_result=model_result,
            max_target_feed_rate_L_h=max_target_feed if max_target_feed > 0 else None,
            max_biomass_g_L=max_biomass if max_biomass > 0 else None,
        )
    except ValueError as exc:
        st.error(f"Phase 3 assessment could not be completed: {exc}")
        return None

    st.session_state["phase3_assessment"] = assessment

    summary = assessment["summary"]
    s1, s2, s3, s4 = st.columns(4)
    with s1:
        st.metric("Engineering checks evaluated", summary["engineering_checks_evaluated"])
    with s2:
        st.metric("Constraints evaluated", summary["constraints_evaluated"])
    with s3:
        st.metric("Checks requiring data", summary["unresolved"])
    with s4:
        st.metric("Deviations / exceedances", summary["deviations"])

    if assessment["large_scale_transition"]:
        st.warning(
            f"Large scale transition: {assessment['scale_factor']:.4g}×. "
            "The assessment does not establish scale-up equivalence or equipment suitability."
        )

    st.markdown("#### Engineering criteria comparison")
    criteria_rows = []
    for item in assessment["criteria"]:
        status_labels = {
            "meets_basis": "Meets selected basis",
            "deviates_from_basis": "Deviates from selected basis",
            "screened_match": "Screened match",
            "different": "Different",
            "meets_constraint": "Within constraint",
            "within_constraint": "Within constraint",
            "at_constraint": "At constraint limit",
            "exceeds_constraint": "Exceeds constraint",
            "capacity_available": "Within capacity",
            "capacity_check": "Exceeds capacity",
            "otr_calculated_our_not_evaluated": "OTR calculated; OUR not evaluated",
            "not_evaluated": "Not evaluated",
        }
        if item.get("category") == "constraint":
            comparison = (
                "Not evaluated"
                if item.get("utilisation_percent") is None
                else f"{item['utilisation_percent']:.3g}% of limit"
            )
        else:
            comparison = (
                "Not evaluated"
                if item.get("ratio") is None
                else f"{item['ratio']:.6g}×"
            )
        criteria_rows.append(
            {
                "Category": "Constraint" if item.get("category") == "constraint" else "Engineering check",
                "Criterion": item["criterion"],
                "Laboratory": "Not evaluated" if item["laboratory"] is None else f"{item['laboratory']:.6g}",
                "Target / actual": "Not evaluated" if item["target"] is None else f"{item['target']:.6g}",
                "Unit": item["unit"],
                "Comparison": comparison,
                "Status": status_labels.get(item["status"], item["status"]),
            }
        )
    st.dataframe(criteria_rows, width="stretch", hide_index=True)

    st.markdown("#### Constraint and data review")
    if assessment["deviations"]:
        st.warning("The following items require engineering review or exceed a supplied constraint.")
        for item in assessment["deviations"]:
            if item.get("category") == "constraint" and item.get("utilisation_percent") is not None:
                st.write(
                    f"- **{item['criterion']}** — {item['utilisation_percent']:.3g}% of limit; "
                    f"**{item['status']}**"
                )
            else:
                st.write(f"- **{item['criterion']}** — {item['status']}")
    else:
        st.success("No supplied constraint was exceeded and no engineering-basis deviation was identified.")

    if assessment["unresolved"]:
        st.info("The following items cannot be fully assessed with the available data:")
        for item in assessment["unresolved"]:
            st.write(f"- **{item['criterion']}** — {item['status']}")

    if assessment.get("data_requirements"):
        st.markdown("#### Data required to complete the assessment")
        for requirement in assessment["data_requirements"]:
            st.write(f"- **{requirement['criterion']}** — {requirement['required_data']}")

    st.markdown("#### Engineering limitations")
    for limitation in assessment["limitations"]:
        st.write(f"- {limitation}")

    with st.expander("Show Phase 3 assessment logic"):
        st.code(
            "Engineering-check ratio = target / laboratory\n"
            "Engineering deviation (%) = (target/laboratory − 1) × 100\n"
            "Constraint utilisation (%) = actual target value / maximum allowed value × 100\n"
            "Constraint status: within limit, at limit, or exceeds limit\n"
            "Vessel capacity = expected final volume compared with stated maximum\n"
            "OTR = kLa × (C* − CL); adequacy requires an independent OUR assessment\n"
            "Heat-transfer capacity = U × A × ΔT compared with stated heat duty",
            language="text",
        )

    return assessment


# -------------------------------------------------------------------
# Streamlit application
# -------------------------------------------------------------------

st.set_page_config(
    page_title="ScaleUp AI",
    page_icon="🧪",
    layout="wide",
)

st.title("ScaleUp AI — Scientific Process Intelligence")
st.caption(
    "Turn process descriptions into structured, traceable process knowledge, "
    "then move from understanding to models, simulation and engineering scale-up."
)

with st.expander("How ScaleUp AI works", expanded=False):
    st.markdown(
        """
        **1. Understand** — extract only what the supplied process description supports.

        **2. Graph** — represent explicit operations and relationships so the process topology can be inspected.

        **3. Evidence** — show where each important claim came from. For uploaded PDFs, evidence can retain the source page.

        **4. Uncertainty** — keep missing, ambiguous, and contradictory information visible instead of silently filling gaps.

        **5. Model** — determine which scientific models are applicable and whether the available inputs are sufficient to run them.

        **6. Simulate** — execute an applicable scientific model only when its required inputs are available.

        **7. Scale** — apply engineering scale-up calculations to the model result, with assumptions and alternative criteria shown explicitly.

        **8. Export** — preserve the structured analysis and downstream calculations for review.
        """
    )

with st.sidebar:
    st.header("Example input")
    st.selectbox("Example", list(EXAMPLES), key="example_name")
    st.button("Load selected example", on_click=load_selected_example)

    st.markdown("---")

    st.caption("ScaleUp AI workflow")
    st.markdown(
        """
        1. Understand — structured ProcessState
        2. Graph — explicit process topology
        3. Evidence — claim/source traceability
        4. Uncertainty — missing / ambiguous / contradictory
        5. Model — applicable scientific models
        6. Simulate — scientific model execution
        7. Scale — engineering scale-up
        8. Export — structured analysis
        """
    )

with st.form("process_description_form"):
    uploaded_pdf = st.file_uploader(
        "Optional scientific PDF",
        type=["pdf"],
        help=(
            "Upload a text-based PDF to extract page-aware evidence. "
            "For scanned/image-only PDFs, OCR is not yet included."
        ),
    )

    source_mode = st.radio(
        "Analysis source",
        ["Process description", "Scientific PDF", "Both — description primary"],
        horizontal=True,
        help=(
            "Process description uses the editable text. Scientific PDF uses the uploaded document as the primary source. "
            "Both keeps the description as the primary process definition while retaining the PDF as supporting document context."
        ),
    )

    description = st.text_area(
        "Process description",
        key="process_description",
        height=300,
        placeholder="Paste a process description here, or upload a PDF above...",
    )

    submitted = st.form_submit_button("Understand process")


if submitted:
    if source_mode == "Scientific PDF":
        if uploaded_pdf is None:
            st.error("Select a scientific PDF before using Scientific PDF as the analysis source.")
            st.stop()
        process_state = understand_pdf(
            uploaded_pdf.getvalue(),
            source_name=uploaded_pdf.name,
        )
        st.session_state["input_source"] = uploaded_pdf.name
    elif source_mode == "Both — description primary":
        process_state = understand_process(description)
        if uploaded_pdf is not None:
            process_state = attach_document_context(
                process_state,
                uploaded_pdf.getvalue(),
                uploaded_pdf.name,
            )
        else:
            st.warning("Both mode selected, but no PDF was uploaded. The process description will be used on its own.")
        st.session_state["input_source"] = "description + supporting_pdf" if uploaded_pdf is not None else "pasted_text"
    else:
        process_state = understand_process(description)
        st.session_state["input_source"] = "pasted_text"

    st.session_state["process_state"] = process_state
    st.session_state["model_result"] = None
    st.session_state["model_error"] = None
    st.session_state["model_readiness"] = None
    st.session_state["scaleup_result"] = None
    st.session_state["engineering_scaleup_result"] = None
    st.session_state["phase3_assessment"] = None

    model_applicable = is_ecoli_model_applicable(
        process_state
    )

    st.session_state["model_applicable"] = model_applicable

    if model_applicable:
        readiness = get_ecoli_model_readiness(
            process_state
        )

        st.session_state["model_readiness"] = readiness

        if readiness["ready"]:
            try:
                st.session_state["model_result"] = (
                    run_from_process_state(
                        process_state
                    )
                )
            except (TypeError, ValueError) as exc:
                st.session_state["model_result"] = None
                st.session_state["model_error"] = str(exc)


if state := st.session_state.get("process_state"):
    # The application now follows the product phases while preserving the
    # existing scientific and engineering engines underneath them.
    tabs = st.tabs([
        "Understand",
        "Graph",
        "Evidence",
        "Uncertainty",
        "Model",
        "Simulate",
        "Scale",
        "Export",
    ])

    with tabs[0]:
        display_process_state(state)

    with tabs[1]:
        display_process_graph(state)

    with tabs[2]:
        display_evidence(state)

    with tabs[3]:
        display_uncertainty(state)

    model_applicable = st.session_state.get("model_applicable", False)
    model_readiness = st.session_state.get("model_readiness")
    model_result = st.session_state.get("model_result")
    model_error = st.session_state.get("model_error")

    with tabs[4]:
        display_model_registry(state)
        with st.expander("Current model-readiness record"):
            st.json(model_readiness or {})

    with tabs[5]:
        display_simulation_status(
            state,
            model_applicable,
            model_readiness,
            model_result,
            model_error,
        )

    scaleup_result = None
    engineering_scaleup_result = None
    phase3_assessment = None

    with tabs[6]:
        if model_result:
            scaleup_result = display_scale_up_section(state, model_result)
            display_scale_up_result(scaleup_result)

            if scaleup_result is not None:
                engineering_scaleup_result = display_engineering_scale_up_section(
                    scaleup_result,
                    model_result,
                )

            if scaleup_result is not None and engineering_scaleup_result is not None:
                phase3_assessment = display_phase3_assessment(
                    scaleup_result,
                    engineering_scaleup_result,
                    model_result,
                )
        else:
            st.info(
                "Run an applicable scientific model first. The Scale stage consumes "
                "the model result and then exposes volumetric, engineering, oxygen-transfer "
                "and thermal screening."
            )

    # Keep computed results available across reruns even when the Scale tab is
    # not the active tab.
    if scaleup_result is not None:
        st.session_state["scaleup_result"] = scaleup_result
    if engineering_scaleup_result is not None:
        st.session_state["engineering_scaleup_result"] = engineering_scaleup_result
    if phase3_assessment is not None:
        st.session_state["phase3_assessment"] = phase3_assessment

    with tabs[7]:
        st.subheader("Export")
        st.info("**What this means:** Export preserves the structured process record and any calculations already performed. The JSON export is the machine-readable ProcessState; Excel is the review-oriented tabular export.")
        stored_scaleup = st.session_state.get("scaleup_result")
        stored_engineering = st.session_state.get("engineering_scaleup_result")
        stored_phase3 = st.session_state.get("phase3_assessment")
        excel_file = create_excel_export(
            state,
            model_result,
            stored_scaleup,
            stored_engineering,
            stored_phase3,
        )
        st.download_button(
            label=(
                "Export analysis + scale-up to Excel"
                if stored_scaleup is not None
                else "Export analysis to Excel"
            ),
            data=excel_file,
            file_name="scaleup_ai_analysis.xlsx",
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            key="export_analysis",
        )

        json_payload = {
            "process_state": state.model_dump(mode="json"),
            "model_result": model_result.model_dump(mode="json") if model_result else None,
            "scaleup_result": stored_scaleup,
            "engineering_scaleup_result": stored_engineering,
            "phase3_assessment": stored_phase3,
        }
        st.download_button(
            label="Export machine-readable JSON",
            data=json.dumps(json_payload, indent=2, ensure_ascii=False),
            file_name="scaleup_ai_analysis.json",
            mime="application/json",
            key="export_json",
        )
        st.caption(
            "The export contains the ProcessState, evidence, graph topology, "
            "scientific model results, engineering calculations and Phase 3 assessment "
            "when those layers have been executed."
        )

